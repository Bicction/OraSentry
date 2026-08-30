# -*- coding: utf-8 -*-
"""企业级 Oracle 巡检 DOCX 报告生成器。"""
import datetime
import os
import re
import subprocess
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Tuple

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html as lxml_html

from config import CATEGORY_NAMES, CATEGORY_ORDER, REPORT_CONFIG
from explanations import get_check_explanation
from inspection_summary import InspectionSummary, build_inspection_summary
from parser.base import CheckResult
from scoring import calculate_score_breakdown, score_band


COLORS = {
    "navy": "17365D",
    "blue": "1F4E78",
    "text": "243447",
    "muted": "64748B",
    "header": "EAF0F6",
    "OK": "059669",
    "OK_BG": "ECFDF5",
    "WARN": "D97706",
    "WARN_BG": "FFFBEB",
    "CRIT": "DC2626",
    "CRIT_BG": "FEF2F2",
    "INFO": "0284C7",
    "INFO_BG": "E0F2FE",
    "UNKNOWN": "64748B",
    "UNKNOWN_BG": "F1F5F9",
}
STATUS_LABELS = {
    "OK": "正常", "WARN": "警告", "CRIT": "严重",
    "INFO": "信息", "UNKNOWN": "数据异常",
}
LATIN_FONT = "Times New Roman"
EAST_ASIA_FONT = "SimSun"
THEME_FONT_ATTRS = (
    qn("w:asciiTheme"),
    qn("w:hAnsiTheme"),
    qn("w:eastAsiaTheme"),
    qn("w:cstheme"),
)


def _ordered_categories(results: Dict[str, List[CheckResult]]) -> List[str]:
    return (
        [category for category in CATEGORY_ORDER if category in results]
        + [category for category in results if category not in CATEGORY_ORDER]
    )


def _set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def _apply_bilingual_fonts(fonts) -> None:
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    for attr in THEME_FONT_ATTRS:
        if fonts.get(attr) is not None:
            del fonts.attrib[attr]


def _set_run_font(run, size=None) -> None:
    run.font.name = LATIN_FONT
    if size is not None:
        run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    _apply_bilingual_fonts(fonts)


def _set_simsun_run_font(run, size=None) -> None:
    """封面项目名称强制使用宋体，包括英文字符。"""
    run.font.name = EAST_ASIA_FONT
    run.font.italic = False
    if size is not None:
        run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in (qn("w:ascii"), qn("w:hAnsi"), qn("w:cs"), qn("w:eastAsia")):
        fonts.set(attribute, EAST_ASIA_FONT)
    for attr in THEME_FONT_ATTRS:
        if fonts.get(attr) is not None:
            del fonts.attrib[attr]


def _set_cell_text(cell, text, *, bold=False, color=None, size=9,
                   align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text if text is not None else ""))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    _set_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_style_font(style, size: float, color: str = COLORS["text"],
                    bold=None) -> None:
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    _apply_bilingual_fonts(fonts)


def _set_document_default_fonts(document: Document) -> None:
    styles = document.styles.element
    doc_defaults = styles.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        return
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    _apply_bilingual_fonts(fonts)


def _configure_styles(document: Document) -> None:
    styles = document.styles
    _set_document_default_fonts(document)
    _set_style_font(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.line_spacing = 1.35
    styles["Normal"].paragraph_format.space_after = Pt(5)
    _set_style_font(styles["Title"], 28, COLORS["navy"], True)
    _set_style_font(styles["Subtitle"], 13, COLORS["muted"])
    _set_style_font(styles["Heading 1"], 18, COLORS["navy"], True)
    _set_style_font(styles["Heading 2"], 13, COLORS["blue"], True)
    _set_style_font(styles["Heading 3"], 11, COLORS["blue"], True)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.space_before = Pt(10)
        styles[name].paragraph_format.space_after = Pt(6)
    for extra in ("List Bullet", "List Number"):
        try:
            _set_style_font(styles[extra], 10.5)
        except KeyError:
            continue


def _add_heading(document: Document, text: str, level: int):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        _set_run_font(run)
    return paragraph


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    """给标题添加 Word 内部书签，供静态目录跳转。"""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_internal_link(paragraph, text: str, anchor: str, *, level: int) -> None:
    """插入不依赖外部关系或域更新的 Word 文档内链接。"""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    _apply_bilingual_fonts(fonts)
    run_properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLORS["navy"] if level == 1 else COLORS["text"])
    run_properties.append(color)
    if level == 1:
        run_properties.append(OxmlElement("w:b"))
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21" if level == 1 else "19")
    run_properties.append(size)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _iter_table_paragraphs(table) -> Iterable:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_document_paragraphs(document: Document) -> Iterable:
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        yield from _iter_table_paragraphs(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                yield paragraph
            for table in part.tables:
                yield from _iter_table_paragraphs(table)


def _normalize_document_fonts(document: Document, cover_project_name: str = "") -> None:
    """全文强制中文宋体、英文数字 Times New Roman，去掉主题字体残留。"""
    _set_document_default_fonts(document)
    for style in document.styles:
        rpr = getattr(style._element, "rPr", None)
        if rpr is None:
            continue
        fonts = getattr(rpr, "rFonts", None)
        if fonts is not None:
            _apply_bilingual_fonts(fonts)
    for paragraph in _iter_document_paragraphs(document):
        for run in paragraph.runs:
            if cover_project_name and paragraph.text == cover_project_name:
                _set_simsun_run_font(run)
            else:
                _set_run_font(run)


def _configure_page(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)


def _clear_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for container in (section.header, section.footer):
        for paragraph in container.paragraphs:
            paragraph.text = ""


def _add_field(paragraph, instruction: str, placeholder: str = "", *, dirty=False) -> None:
    run = paragraph.add_run()
    _set_run_font(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def _add_toc_page_reference(paragraph, anchor: str) -> None:
    """给静态目录添加内部书签页码；不是外部链接，也不使用 TOC 域。"""
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    run = paragraph.add_run("\t")
    _set_run_font(run, 9.5)
    _add_field(paragraph, f"PAGEREF {anchor} \\h", "—")


def _cache_toc_page_numbers(output_path: Path) -> bool:
    """用本机 Word 分页引擎写入准确页码并锁定内部 PAGEREF 域。"""
    if os.name != "nt":
        return False
    script = r"""
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($env:ORACLE_REPORT_DOCX_PATH, $false, $false)
    $document.Repaginate()
    foreach ($field in $document.Fields) {
        if ($field.Code.Text -match 'PAGEREF\s+([^\s\\]+)') {
            $anchor = $Matches[1]
            if ($document.Bookmarks.Exists($anchor)) {
                $page = $document.Bookmarks.Item($anchor).Range.Information(1)
                $field.Result.Text = [string]$page
                $field.Locked = $true
            }
        }
    }
    $document.Save()
} finally {
    if ($document -ne $null) { $document.Close($true) }
    if ($word -ne $null) { $word.Quit() }
}
"""
    environment = os.environ.copy()
    environment["ORACLE_REPORT_DOCX_PATH"] = str(output_path.resolve())
    creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    try:
        completed = subprocess.run(
            ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
            env=environment,
            capture_output=True,
            text=True,
            timeout=90,
            creationflags=creation_flags,
            check=False,
        )
        return completed.returncode == 0
    except (OSError, subprocess.SubprocessError):
        return False


def _configure_content_header_footer(section, title: str, hostname: str,
                                     footer_text: str) -> None:
    _clear_header_footer(section)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run(f"{title}  |  {hostname}")
    _set_run_font(header_run, 8.5)
    header_run.font.color.rgb = RGBColor.from_string(COLORS["muted"])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{footer_text}  ·  第 ")
    _set_run_font(run, 8.5)
    run.font.color.rgb = RGBColor.from_string(COLORS["muted"])
    _add_field(footer, "PAGE", "1")
    run = footer.add_run(" 页")
    _set_run_font(run, 8.5)
    run.font.color.rgb = RGBColor.from_string(COLORS["muted"])

    sect_pr = section._sectPr
    page_num = sect_pr.find(qn("w:pgNumType"))
    if page_num is None:
        page_num = OxmlElement("w:pgNumType")
        sect_pr.append(page_num)
    page_num.set(qn("w:start"), "1")


def _configure_cover_section(document: Document) -> None:
    """配置首节，使正式封面直接从第一页开始。"""
    section = document.sections[0]
    _configure_page(section)
    _clear_header_footer(section)


def _format_collect_time(value) -> str:
    text = str(value or "N/A")
    try:
        if len(text) == 15 and "_" in text:
            return datetime.datetime.strptime(text, "%Y%m%d_%H%M%S").strftime(
                "%Y-%m-%d %H:%M:%S"
            )
    except ValueError:
        pass
    return text


def _report_title(results) -> str:
    if any(category in results for category in ("db", "cdb", "rac", "security")):
        return REPORT_CONFIG["title_db"]
    if "host" in results:
        return REPORT_CONFIG["title_host"]
    return REPORT_CONFIG["title_general"]


def _overall_status(counts) -> str:
    if counts["CRIT"]:
        return "存在严重问题"
    if counts["WARN"]:
        return "存在警告"
    return "正常"


def _add_cover(document: Document, results, env_info, health_score, counts) -> str:
    title = _report_title(results)
    score_breakdown = calculate_score_breakdown(results)
    for _ in range(4):
        document.add_paragraph()
    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    _set_run_font(title_run)
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    project_name = str(
        env_info.get("project_name", "") or REPORT_CONFIG["project_name"]
    )
    subtitle_run = subtitle.add_run(project_name)
    _set_simsun_run_font(subtitle_run, 13)

    line = document.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("━" * 32)
    _set_run_font(run)
    run.font.color.rgb = RGBColor.from_string(COLORS["blue"])

    score = document.add_paragraph()
    score.alignment = WD_ALIGN_PARAGRAPH.CENTER
    score_run = score.add_run(f"{health_score}")
    _set_run_font(score_run, 34)
    score_run.bold = True
    score_color = COLORS[score_band(health_score, score_breakdown.possible > 0)]
    score_run.font.color.rgb = RGBColor.from_string(score_color)
    label = score.add_run("\n健康评分 / 100    ")
    _set_run_font(label, 11)
    overall_key = "CRIT" if counts.get("CRIT") else "WARN" if counts.get("WARN") else "OK"
    status_run = score.add_run(_overall_status(counts))
    _set_run_font(status_run, 11)
    status_run.bold = True
    status_run.font.color.rgb = RGBColor.from_string(COLORS[overall_key])
    confidence_run = score.add_run(f"    数据可信度：{score_breakdown.confidence}")
    _set_run_font(confidence_run, 11)

    meta = document.add_table(rows=6, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = True
    values = [
        ("主机名称", env_info.get("hostname", "N/A")),
        ("服务器地址", env_info.get("server_ip", "N/A")),
        ("数据库 SID", env_info.get("oracle_sid", "N/A")),
        ("巡检时间", _format_collect_time(env_info.get("timestamp"))),
        ("数据可信度", score_breakdown.confidence),
        ("报告标识", REPORT_CONFIG["brand"]),
    ]
    for row, (label_text, value) in zip(meta.rows, values):
        _set_cell_shading(row.cells[0], COLORS["header"])
        _set_cell_text(row.cells[0], label_text, bold=True, color=COLORS["navy"], size=10)
        _set_cell_text(row.cells[1], value, size=10)

    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(24)
    run = footer.add_run("企业级数据库运行健康与安全巡检")
    _set_run_font(run, 10)
    run.font.color.rgb = RGBColor.from_string(COLORS["muted"])
    return title


def _add_summary(document: Document, results, counts, health_score) -> None:
    score_breakdown = calculate_score_breakdown(results)
    _add_heading(document, "执行摘要", 1)
    document.add_paragraph(
        f"本次巡检共覆盖 {counts['TOTAL']} 个检查项目，健康评分为 {health_score}/100，"
        f"总体结论：{_overall_status(counts)}；数据可信度：{score_breakdown.confidence}。"
    )

    table = document.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ("正常", "警告", "严重", "信息", "数据异常", "总计")
    values = (
        counts["OK"], counts["WARN"], counts["CRIT"],
        counts.get("INFO", 0), counts.get("UNKNOWN", 0), counts["TOTAL"],
    )
    fills = (
        COLORS["OK_BG"], COLORS["WARN_BG"], COLORS["CRIT_BG"],
        COLORS["INFO_BG"], COLORS["UNKNOWN_BG"], COLORS["header"],
    )
    colors = (
        COLORS["OK"], COLORS["WARN"], COLORS["CRIT"],
        COLORS["INFO"], COLORS["UNKNOWN"], COLORS["navy"],
    )
    for index, label in enumerate(labels):
        _set_cell_shading(table.cell(0, index), fills[index])
        _set_cell_text(table.cell(0, index), label, bold=True, color=colors[index],
                       align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(table.cell(1, index), values[index], bold=True, color=colors[index],
                       size=16, align=WD_ALIGN_PARAGRAPH.CENTER)

    actions = []
    for category in _ordered_categories(results):
        for item in results[category]:
            if item.status in ("UNKNOWN", "CRIT", "WARN"):
                actions.append((category, item))
    actions.sort(key=lambda row: {"UNKNOWN": 0, "CRIT": 1, "WARN": 2}[row[1].status])
    _add_heading(document, "待处理事项", 2)
    if not actions:
        document.add_paragraph("未发现需要跟进的警告或严重问题。")
    for category, item in actions:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(
            f"[{STATUS_LABELS[item.status]}] "
            f"{CATEGORY_NAMES.get(category, category)} / {item.name}：{item.value}"
        )
        _set_run_font(run, 9.5)
        run.font.color.rgb = RGBColor.from_string(COLORS[item.status])
        if item.suggestion:
            suggestion = paragraph.add_run(f"；{item.suggestion}")
            _set_run_font(suggestion, 9.5)


def _add_toc(document: Document, results: Dict[str, List[CheckResult]]) -> None:
    """添加打开即显示的静态目录；不使用 TOC 域，避免 Word 更新域提示。"""
    document.add_page_break()
    heading = _add_heading(document, "目录", 1)
    _add_bookmark(heading, "report_toc", 1)

    note = document.add_paragraph("点击目录项可直接跳转到对应巡检章节。")
    note.runs[0].font.color.rgb = RGBColor.from_string(COLORS["muted"])

    categories = _ordered_categories(results)
    for category_index, category in enumerate(categories, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(2)
        _add_internal_link(
            paragraph,
            f"{category_index}. {CATEGORY_NAMES.get(category, category)}",
            f"report_category_{category_index}",
            level=1,
        )
        _add_toc_page_reference(paragraph, f"report_category_{category_index}")
        for item_index, item in enumerate(results[category], start=1):
            item_paragraph = document.add_paragraph()
            item_paragraph.paragraph_format.left_indent = Cm(0.8)
            item_paragraph.paragraph_format.space_after = Pt(1)
            _add_internal_link(
                item_paragraph,
                f"{category_index}.{item_index}  {item.name}",
                f"report_item_{category_index}_{item_index}",
                level=2,
            )
            _add_toc_page_reference(
                item_paragraph, f"report_item_{category_index}_{item_index}"
            )

    summary_index = len(categories) + 1
    summary_entry = document.add_paragraph()
    summary_entry.paragraph_format.space_before = Pt(4)
    _add_internal_link(
        summary_entry,
        f"{summary_index}. 总结",
        "report_summary",
        level=1,
    )
    _add_toc_page_reference(summary_entry, "report_summary")


def _normalize_text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _class_names(element) -> set:
    return set((element.get("class") or "").split())


def _with_class(element, class_name: str):
    return element.xpath(
        f'.//*[contains(concat(" ", normalize-space(@class), " "), " {class_name} ")]'
    )


def _table_block(element):
    rows = []
    for tr in element.xpath(".//tr"):
        cells = [_normalize_text(cell.text_content()) for cell in tr.xpath("./th|./td")]
        if cells:
            rows.append(cells)
    return ("table", rows) if rows else None


def _bar_blocks(element) -> List[Tuple]:
    blocks = []
    for row in _with_class(element, "bar-row"):
        label_nodes = _with_class(row, "bar-label")
        value_nodes = _with_class(row, "bar-value")
        fill_nodes = _with_class(row, "bar-fill")
        label = _normalize_text(label_nodes[0].text_content()) if label_nodes else "指标"
        value = _normalize_text(value_nodes[0].text_content()) if value_nodes else ""
        status = "OK"
        if fill_nodes:
            classes = _class_names(fill_nodes[0])
            status = "CRIT" if "bar-fill-crit" in classes else "WARN" if "bar-fill-warn" in classes else "OK"
        blocks.append(("bar", label, value, status))
    return blocks


def _health_block(element):
    head_nodes = _with_class(element, "health-percentage-head")
    head = head_nodes[0] if head_nodes else element
    label_nodes = head.xpath(".//span")
    value_nodes = head.xpath(".//strong")
    fill_nodes = _with_class(element, "health-percentage-fill")
    label = _normalize_text(label_nodes[0].text_content()) if label_nodes else "命中率"
    value = _normalize_text(value_nodes[0].text_content()) if value_nodes else ""
    status = "OK"
    if fill_nodes:
        classes = _class_names(fill_nodes[0])
        status = "CRIT" if "bar-fill-crit" in classes else "WARN" if "bar-fill-warn" in classes else "OK"
    return ("bar", label, value, status)


def _element_blocks(element) -> Iterable[Tuple]:
    tag = str(element.tag).lower()
    classes = _class_names(element)
    if tag == "table":
        block = _table_block(element)
        if block:
            yield block
        return
    if "bar-chart" in classes:
        yield from _bar_blocks(element)
        return
    if "health-percentage" in classes:
        yield _health_block(element)
        return
    if "detail-subtitle" in classes:
        text = _normalize_text(element.text_content())
        if text:
            yield ("subtitle", text)
        return
    if "detail-note" in classes:
        text = _normalize_text(element.text_content())
        if text:
            yield ("note", text)
        return
    if tag in ("ul", "ol"):
        items = [
            _normalize_text(li.text_content())
            for li in element.xpath("./li")
            if _normalize_text(li.text_content())
        ]
        if items:
            yield ("bullets", items)
        return

    children = list(element)
    if children:
        emitted = False
        for child in children:
            for block in _element_blocks(child):
                emitted = True
                yield block
        if not emitted:
            text = _normalize_text(element.text_content())
            if text:
                yield ("paragraph", text)
        return

    text = _normalize_text(element.text_content())
    if text:
        yield ("paragraph", text)


def parse_extra_html(extra_html: str) -> List[Tuple]:
    """把报告中的受控 HTML 片段转换为顺序化 Word 内容块。"""
    if not extra_html:
        return []
    try:
        root = lxml_html.fragment_fromstring(extra_html, create_parent="div")
        blocks = []
        for child in root:
            blocks.extend(_element_blocks(child))
        if blocks:
            return blocks
        text = _normalize_text(root.text_content())
        return [("paragraph", text)] if text else []
    except (ValueError, TypeError):
        text = _normalize_text(re.sub(r"<[^>]+>", " ", extra_html))
        return [("paragraph", text)] if text else []


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_word_table(document: Document, rows: List[List[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    font_size = 7.5 if width >= 7 else 8.5 if width >= 5 else 9
    for row_index, values in enumerate(normalized):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            if row_index == 0:
                _set_cell_shading(cell, COLORS["header"])
            _set_cell_text(
                cell, value, bold=row_index == 0,
                color=COLORS["navy"] if row_index == 0 else COLORS["text"],
                size=font_size,
            )
    _set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_visual_bar(document: Document, label: str, value: str, status: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    label_run = paragraph.add_run(f"{label:<22} ")
    _set_run_font(label_run, 9)
    match = re.search(r"-?\d+(?:\.\d+)?", value)
    numeric = max(0.0, min(float(match.group()) if match else 0.0, 100.0))
    filled = max(1, round(numeric / 5)) if numeric > 0 else 0
    bar_run = paragraph.add_run("■" * filled)
    _set_run_font(bar_run, 9)
    bar_run.font.color.rgb = RGBColor.from_string(COLORS.get(status, COLORS["OK"]))
    empty_run = paragraph.add_run("□" * (20 - filled))
    _set_run_font(empty_run, 9)
    empty_run.font.color.rgb = RGBColor.from_string("CBD5E1")
    value_run = paragraph.add_run(f"  {value}")
    _set_run_font(value_run, 9)
    value_run.bold = True


def _render_extra_blocks(document: Document, extra_html: str) -> None:
    for block in parse_extra_html(extra_html):
        kind = block[0]
        if kind == "table":
            _add_word_table(document, block[1])
        elif kind == "subtitle":
            _add_heading(document, block[1], 3)
        elif kind == "note":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block[1])
            _set_run_font(run, 8.5)
            run.italic = True
            run.font.color.rgb = RGBColor.from_string(COLORS["muted"])
        elif kind == "bullets":
            for text in block[1]:
                document.add_paragraph(text, style="List Bullet")
        elif kind == "bar":
            _add_visual_bar(document, block[1], block[2], block[3])
        elif kind == "paragraph":
            document.add_paragraph(block[1])


def _add_status_table(document: Document, item: CheckResult) -> None:
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    labels = ("状态", "当前值")
    values = (STATUS_LABELS.get(item.status, item.status), item.value)
    for index, label in enumerate(labels):
        _set_cell_shading(table.cell(0, index), COLORS["header"])
        _set_cell_text(table.cell(0, index), label, bold=True, color=COLORS["navy"])
        fill = COLORS.get(f"{item.status}_BG", COLORS["header"]) if index == 0 else "FFFFFF"
        color = COLORS.get(item.status, COLORS["text"]) if index == 0 else COLORS["text"]
        _set_cell_shading(table.cell(1, index), fill)
        _set_cell_text(table.cell(1, index), values[index], bold=index == 0, color=color, size=10)


def _add_check_item(document: Document, category_index: int, item_index: int,
                    item: CheckResult) -> None:
    heading = _add_heading(document, f"{category_index}.{item_index}  {item.name}", 2)
    _add_bookmark(
        heading,
        f"report_item_{category_index}_{item_index}",
        10000 + category_index * 1000 + item_index,
    )
    _add_status_table(document, item)
    _add_heading(document, "巡检项说明", 3)
    document.add_paragraph(get_check_explanation(item.name))
    _add_heading(document, "检查结果", 3)
    document.add_paragraph(item.detail or "无补充说明。")
    if item.suggestion:
        _add_heading(document, "整改建议", 3)
        table = document.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        _set_cell_shading(table.cell(0, 0), COLORS["WARN_BG"])
        _set_cell_text(table.cell(0, 0), item.suggestion, color=COLORS["WARN"], size=9.5)
    if item.extra_html:
        _add_heading(document, "巡检明细", 3)
        _render_extra_blocks(document, item.extra_html)


def _add_report_summary_chapter(document: Document, section_index: int,
                                content: InspectionSummary = None) -> None:
    """添加总结章节；未要求生成内容时保留可填写的空白区域。"""
    heading = _add_heading(document, f"{section_index}. 总结", 1)
    heading.paragraph_format.page_break_before = True
    _add_bookmark(heading, "report_summary", 900000)
    if content is None:
        for _ in range(6):
            document.add_paragraph()
        return

    _add_heading(document, "总体结论", 2)
    document.add_paragraph(content.overview)
    document.add_paragraph(content.conclusion)

    _add_heading(document, "重点问题", 2)
    if content.issues:
        for issue in content.issues:
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(issue)
            _set_run_font(run, 10)
            if issue.startswith("[严重]"):
                run.font.color.rgb = RGBColor.from_string(COLORS["CRIT"])
            elif issue.startswith("[警告]"):
                run.font.color.rgb = RGBColor.from_string(COLORS["WARN"])
        if content.omitted_issues:
            note = document.add_paragraph(
                f"另有 {content.omitted_issues} 项风险未在本章展开，请结合待处理事项和巡检明细查看。"
            )
            note.runs[0].font.color.rgb = RGBColor.from_string(COLORS["muted"])
    else:
        document.add_paragraph("未发现警告或严重问题。")

    _add_heading(document, "处置建议", 2)
    if content.recommendations:
        for recommendation in content.recommendations:
            paragraph = document.add_paragraph(style="List Number")
            run = paragraph.add_run(recommendation)
            _set_run_font(run, 10)
        if content.omitted_recommendations:
            note = document.add_paragraph(
                f"另有 {content.omitted_recommendations} 条建议未在本章展开，请查阅对应巡检项。"
            )
            note.runs[0].font.color.rgb = RGBColor.from_string(COLORS["muted"])
    else:
        document.add_paragraph("继续保持现有运维策略，并按计划开展周期性复查。")


def generate_docx_report(results: Dict[str, List[CheckResult]], env_info: dict,
                         output_file: str, health_score: int, counts: dict,
                         generate_summary_content: bool = False,
                         cancellation_check: Optional[Callable[[], None]] = None) -> int:
    """生成完整 DOCX 报告，返回实际渲染的巡检项数量。"""
    if cancellation_check:
        cancellation_check()
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_styles(document)
    _configure_cover_section(document)
    project_name = str(
        env_info.get("project_name", "") or REPORT_CONFIG["project_name"]
    ).strip() or REPORT_CONFIG["project_name"]
    title = _add_cover(document, results, env_info, health_score, counts)

    content_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    _configure_page(content_section)
    _configure_content_header_footer(
        content_section,
        title,
        str(env_info.get("hostname", "N/A")),
        project_name,
    )
    _add_summary(document, results, counts, health_score)
    _add_toc(document, results)
    if cancellation_check:
        cancellation_check()

    rendered_count = 0
    categories = _ordered_categories(results)
    for category_index, category in enumerate(categories, start=1):
        if cancellation_check:
            cancellation_check()
        document.add_page_break()
        category_name = CATEGORY_NAMES.get(category, category)
        heading = _add_heading(document, f"{category_index}. {category_name}", 1)
        _add_bookmark(
            heading,
            f"report_category_{category_index}",
            1000 + category_index,
        )
        intro = document.add_paragraph(f"本章节共包含 {len(results[category])} 个巡检项目。")
        intro.runs[0].font.color.rgb = RGBColor.from_string(COLORS["muted"])
        for item_index, item in enumerate(results[category], start=1):
            if cancellation_check:
                cancellation_check()
            _add_check_item(document, category_index, item_index, item)
            rendered_count += 1

    if cancellation_check:
        cancellation_check()
    summary_content = (
        build_inspection_summary(results, counts, health_score)
        if generate_summary_content else None
    )
    _add_report_summary_chapter(document, len(categories) + 1, summary_content)

    document.core_properties.title = title
    document.core_properties.subject = "Oracle 数据库与主机巡检报告"
    document.core_properties.author = REPORT_CONFIG["brand"]
    document.core_properties.keywords = "Oracle, Inspection, Database, Security"
    _normalize_document_fonts(document, project_name)
    if cancellation_check:
        cancellation_check()
    document.save(str(output_path))
    _cache_toc_page_numbers(output_path)
    return rendered_count
