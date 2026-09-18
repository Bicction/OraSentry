from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from docx import Document
from docx_gen import _add_word_table, _add_check_item, _normalize_document_fonts, generate_docx_report
from parser.base import CheckResult, generate_data_table


class DocxPerformanceTests(unittest.TestCase):
    def test_long_table_keeps_every_cell_and_special_text(self):
        rows=[['编号','内容']]+[[str(n),f'中文 {n}\nA&B\t<xml>'] for n in range(1200)]
        document=Document()
        _add_word_table(document, rows)
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'table.docx';document.save(path)
            restored=Document(path)
            actual=[[cell.text for cell in row.cells] for row in restored.tables[0].rows]
        self.assertEqual(actual,rows)
        self.assertIn('tblHeader',restored.tables[0].rows[0]._tr.xml)

    def test_cancel_reaches_inside_check_item_table(self):
        class Cancelled(Exception): pass
        calls=[]
        def cancel():
            calls.append(1)
            if len(calls)==12: raise Cancelled()
        item=CheckResult('长列表','OK','1000条','',extra_html=generate_data_table(['ID'],[[n] for n in range(1000)]))
        with self.assertRaises(Cancelled):
            _add_check_item(Document(),1,1,item,cancel)
        self.assertEqual(len(calls),12)

    def test_default_export_does_not_wait_for_word_application(self):
        with tempfile.TemporaryDirectory() as td, patch('docx_gen._cache_toc_page_numbers') as paginate:
            path=Path(td)/'report.docx'
            generate_docx_report({'db':[CheckResult('实例','OK','OPEN','正常')]},{},str(path),100,{'OK':1,'WARN':0,'CRIT':0,'UNKNOWN':0,'INFO':0,'TOTAL':1})
            paginate.assert_not_called()
            self.assertTrue(path.is_file())


if __name__=='__main__': unittest.main()
