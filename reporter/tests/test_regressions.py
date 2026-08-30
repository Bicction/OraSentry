# -*- coding: utf-8 -*-
"""关键误判与安全回归测试。"""
import io
import os
import re
import sys
import tarfile
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = ROOT / "src"
sys.path.insert(0, str(SOURCE_ROOT))

from config import REPORT_CONFIG
from parser.base import CheckResult, generate_data_table, parse_collection_integrity
from parser.db_parser import (
    _parse_buffer_cache_hit,
    _parse_control_files,
    _parse_data_files,
    _parse_database_resilience,
    _parse_dead_processes,
    _parse_instance_status,
    _parse_asm_diskgroups,
    _parse_library_cache_hit,
    _parse_awr_metrics,
    _parse_alert_log,
    _parse_pipe_table,
    _parse_rollback_segments,
    _parse_tablespaces,
    _parse_top_disk_read_sql,
)
from parser.host_parser import (
    _parse_disk_usage,
    _parse_firewall,
    _parse_memory,
    _parse_oracle_host_readiness,
    _parse_time_sync,
)
from parser.security_parser import (
    _parse_audit_and_access_controls,
    _parse_listener_security,
)
from docx_gen import COLORS as DOCX_COLORS, generate_docx_report, parse_extra_html
from gui import (
    add_unique_paths,
    choose_window_geometry,
    classify_input,
    filter_dropped_paths,
    normalize_output_selection,
    score_tone,
)
from report_gen import (
    APP_VERSION,
    ReportBuildError,
    _make_anchor,
    build_report,
    build_reports,
    calculate_health_score,
    default_output_path,
    expand_input_paths,
    extract_tar_gz,
    generate_action_panel,
    generate_category_section,
    generate_report,
    generate_sidebar_nav,
    group_report_jobs,
    project_root,
    report_basename,
    resolve_output_file,
    template_path,
)
from summary_gen import (
    IssueRef,
    SystemSummary,
    build_summary_model,
    generate_summary_docx,
    generate_summary_html,
)
from scoring import calculate_score_breakdown, score_band


class RegressionTests(unittest.TestCase):
    @staticmethod
    def _write(directory, filename, content):
        Path(directory, filename).write_text(content, encoding="utf-8")

    def test_pipe_parser_preserves_null_columns(self):
        rows = _parse_pipe_table("A|B|C\n1||3\n")
        self.assertEqual(rows[1], ["1", "", "3"])

    def test_v43_reporter_and_windows_version_are_consistent(self):
        self.assertEqual(APP_VERSION, "4.3.0")
        version_info = Path(ROOT, "tools", "version_info.txt").read_text(encoding="utf-8")
        self.assertIn("filevers=(4, 3, 0, 0)", version_info)
        self.assertIn("ProductVersion', '4.3.0.0'", version_info)

    def test_collection_integrity_detects_sqlplus_error(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "env.info").write_text("hostname=test\n", encoding="utf-8")
            Path(td, "bad.txt").write_text("ORA-00904: invalid identifier\n", encoding="utf-8")
            result = parse_collection_integrity(td)
            self.assertEqual(result.status, "UNKNOWN")

    def test_collection_integrity_accepts_version_skips(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "collection_manifest.tsv").write_text(
                "item\ttype\tstatus\texit_code\tmessage\n"
                "pdb_info.txt\tSQL\tSKIPPED\t0\t非CDB不适用\n",
                encoding="utf-8",
            )
            result = parse_collection_integrity(td)
            self.assertEqual(result.status, "INFO")
            self.assertIn("1项不适用", result.value)

    def test_collection_integrity_deduplicates_manifest_and_file_error(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "collection_manifest.tsv").write_text(
                "item\ttype\tstatus\texit_code\tmessage\n"
                "bad.txt\tSQL\tFAILED\t136\tORA-00904\n",
                encoding="utf-8",
            )
            Path(td, "bad.txt").write_text("ORA-00904: invalid identifier\n", encoding="utf-8")
            result = parse_collection_integrity(td)
            self.assertEqual(result.value, "1项失败")

    def test_collection_integrity_ignores_collected_alert_messages(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "collection_manifest.tsv").write_text(
                "item\ttype\tstatus\texit_code\tmessage\n",
                encoding="utf-8",
            )
            Path(td, "alert_log_tail.txt").write_text(
                "ORA-16191: Primary log shipping client not logged on standby\n",
                encoding="utf-8",
            )
            result = parse_collection_integrity(td)
            self.assertEqual(result.status, "INFO")

    def test_alert_log_uses_true_count_and_displays_instance_source(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "alert_log_summary.txt",
                "DATABASE_NAME|INSTANCE_NAME|HOST_NAME|ALERT_FILE|ALERT_COUNT|SEVERE_COUNT|FIRST_TIME|LAST_TIME|WINDOW_DAYS|ANALYZED_LINE_LIMIT|ANALYZED_LINE_COUNT|WINDOW_TRUNCATED|ALERT_FILE_SIZE_BYTES|ALERT_FILE_SIZE_DISPLAY|ANALYSIS_START_TIME|ANALYSIS_END_TIME\n"
                "ORCL|orcl1|db01|/u01/app/oracle/diag/rdbms/orcl/orcl1/trace/alert_orcl1.log|1440|0|2026-07-26T10:00:00+08:00|2026-08-24T10:00:00+08:00|30|100000|100000|YES|3221225472|3.00 GB|2026-08-01 08:00:00|2026-08-24 10:00:00\n",
            )
            self._write(
                td, "alert_log_recent.txt",
                "EVENT_TIME|MESSAGE_TEXT\n"
                "2026-08-24T10:00:00+08:00|returning error ORA-16191\n",
            )
            result = _parse_alert_log(td)
            self.assertEqual(result.status, "CRIT")
            self.assertEqual(result.value, "1440条近30天告警 / 日志 3.00 GB")
            self.assertIn("物理日志告警明细: 1 条", result.detail)
            self.assertIn("分析时间范围: 2026-08-01 08:00:00 至 2026-08-24 10:00:00", result.detail)
            self.assertNotIn("文件尾部", result.detail)
            self.assertNotIn("10万行", result.extra_html)
            self.assertIn("超过2GB", result.suggestion)
            self.assertIn("orcl1", result.extra_html)
            self.assertIn("db01", result.extra_html)
            self.assertIn("alert_orcl1.log", result.extra_html)
            self.assertIn("ORA-16191", result.extra_html)

    def test_alert_log_over_2gb_warns_even_without_errors(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "alert_log_summary.txt",
                "DATABASE_NAME|INSTANCE_NAME|HOST_NAME|ALERT_FILE|ALERT_COUNT|SEVERE_COUNT|FIRST_TIME|LAST_TIME|WINDOW_DAYS|ANALYZED_LINE_LIMIT|ANALYZED_LINE_COUNT|WINDOW_TRUNCATED|ALERT_FILE_SIZE_BYTES|ALERT_FILE_SIZE_DISPLAY\n"
                "ORCL|orcl1|db01|/diag/alert_orcl1.log|0|0|||30|100000|85000|NO|2147483649|2.00 GB\n",
            )
            self._write(td, "alert_log_recent.txt", "EVENT_TIME|MESSAGE_TEXT\n")
            result = _parse_alert_log(td)
            self.assertEqual(result.status, "WARN")
            self.assertIn("超过2GB", result.suggestion)
            self.assertIn("物理日志大小", result.extra_html)

    def test_alert_log_marks_legacy_200_row_result_as_truncated(self):
        with tempfile.TemporaryDirectory() as td:
            rows = [
                f"2026-08-24 10:{index // 60:02d}:{index % 60:02d} +08:00|1|16||returning error ORA-16191"
                for index in range(200)
            ]
            self._write(
                td, "alert_log_recent.txt",
                "EVENT_TIME|MESSAGE_TYPE|MESSAGE_LEVEL|PROBLEM_KEY|MESSAGE_TEXT\n"
                + "\n".join(rows) + "\n",
            )
            result = _parse_alert_log(td)
            self.assertEqual(result.value, "至少200条近7天告警")
            self.assertIn("至少200", result.detail)

    def test_boot_efi_disk_is_not_alerted(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "disk_usage.txt",
                "Filesystem      Size  Used Avail Use% Mounted on\n"
                "/dev/sda2        50G   40G  10G  80% /\n"
                "/dev/sda1       200M  199M  1.0G  99% /boot/efi\n"
                "/dev/sda3       500M  400M  100M  80% /boot\n",
            )
            results = _parse_disk_usage(td)
            names = [item.name for item in results]
            self.assertNotIn("磁盘使用率-/boot/efi", names)
            self.assertNotIn("磁盘使用率-/boot", names)
            self.assertTrue(any(item.name == "磁盘使用率-/" for item in results))
            efi_alerts = [
                item for item in results
                if "boot" in item.name.lower() or "/boot/efi" in (item.suggestion or "")
            ]
            self.assertEqual(efi_alerts, [])

    def test_read_only_iso_disk_is_not_alerted(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "disk_usage.txt",
                "Filesystem      Size  Used Avail Use% Mounted on\n"
                "/dev/sda2        50G   20G   30G  40% /\n"
                "/dev/loop0      4.6G  4.6G     0 100% /mnt/ISO\n",
            )
            self._write(
                td, "mounts.txt",
                "/dev/sda2 xfs rw,relatime /\n"
                "/dev/loop0 iso9660 ro,relatime,blocksize=2048 /mnt/ISO\n",
            )

            results = _parse_disk_usage(td)
            names = [item.name for item in results]
            self.assertNotIn("磁盘使用率-/mnt/ISO", names)
            self.assertTrue(any(item.name == "磁盘使用率-/" for item in results))
            overview = next(item for item in results if item.name == "磁盘使用概览")
            self.assertEqual(overview.value, "1个挂载点")
            self.assertIn("已忽略 1 个只读光盘/镜像文件系统", overview.detail)
            self.assertIn("/mnt/ISO(iso9660)", overview.detail)

    def test_writable_disk_under_mnt_is_still_alerted(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "disk_usage.txt",
                "Filesystem      Size  Used Avail Use% Mounted on\n"
                "/dev/sdb1       100G   96G    4G  96% /mnt/data\n",
            )
            self._write(
                td, "mounts.txt",
                "/dev/sdb1 ext4 rw,relatime /mnt/data\n",
            )

            result = next(
                item for item in _parse_disk_usage(td)
                if item.name == "磁盘使用率-/mnt/data"
            )
            self.assertEqual(result.status, "CRIT")

    def test_swap_21_percent_without_activity_is_ok(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "memory.txt",
                "              total        used        free      shared  buff/cache   available\n"
                "Mem:         128354        6169        2464       47962      119720       73110\n"
                "Swap:         12211        2560        9651\n",
            )
            self._write(
                td, "vmstat.txt",
                "procs -----------memory---------- ---swap-- -----io---- -system-- ------cpu-----\n"
                " r  b   swpd   free   buff  cache   si   so    bi    bo   in   cs us sy id wa st\n"
                " 0  0 2622164 2523804 200920 122393120  10   20   268   181    0    0  1  1 98  0  0\n"
                " 0  0 2622164 2523156 200920 122393192   0    0    81    64 10963 16143  1  1 99  0  0\n"
                " 1  0 2622164 2523832 200920 122393192   0    0    65     7 10349 16369  0  0 99  0  0\n",
            )

            result = next(item for item in _parse_memory(td) if item.name == "Swap使用率")
            self.assertEqual(result.status, "OK")
            self.assertEqual(result.value, "21.0%")
            self.assertIn("持续换页: 否", result.detail)
            self.assertEqual(result.suggestion, "")

    def test_memory_prefers_memavailable_and_aligns_displayed_used(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "memory.txt",
                "              total        used        free      shared  buff/cache   available\n"
                "Mem:          64000       10000        5000           0       49000       40000\n"
                "Swap:         10000        2000        8000\n",
            )
            self._write(
                td, "meminfo.txt",
                "MemTotal:       65536000 kB\n"
                "MemFree:         5120000 kB\n"
                "MemAvailable:   40960000 kB\n"
                "Buffers:         1024000 kB\n"
                "Cached:         47104000 kB\n"
                "SReclaimable:      51200 kB\n"
                "Shmem:                 0 kB\n",
            )

            result = next(item for item in _parse_memory(td) if item.name == "内存使用率")
            self.assertEqual(result.value, "37.5%")
            self.assertIn("实际占用: 24000MB", result.detail)
            self.assertIn("可用: 40000MB", result.detail)
            self.assertIn("计算口径: MemAvailable", result.detail)

    def test_centos6_without_memavailable_uses_legacy_meminfo_estimate(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "memory.txt",
                "             total       used       free     shared    buffers     cached\n"
                "Mem:          1000        900        100        100         50        300\n"
                "-/+ buffers/cache:        550        450\n"
                "Swap:          500        100        400\n",
            )
            self._write(
                td, "meminfo.txt",
                "MemTotal:       1024000 kB\n"
                "MemFree:         102400 kB\n"
                "Buffers:          51200 kB\n"
                "Cached:          307200 kB\n"
                "SReclaimable:     51200 kB\n"
                "Shmem:           102400 kB\n",
            )

            result = next(item for item in _parse_memory(td) if item.name == "内存使用率")
            self.assertEqual(result.value, "60.0%")
            self.assertIn("实际占用: 600MB", result.detail)
            self.assertIn("可用: 400MB", result.detail)
            self.assertIn("计算口径: 旧内核兼容估算", result.detail)

    def test_centos6_free_output_falls_back_to_buffers_cache_row(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "memory.txt",
                "             total       used       free     shared    buffers     cached\n"
                "Mem:          1000        900        100        100         50        300\n"
                "-/+ buffers/cache:        550        450\n"
                "Swap:          500        100        400\n",
            )

            result = next(item for item in _parse_memory(td) if item.name == "内存使用率")
            self.assertEqual(result.value, "55.0%")
            self.assertIn("实际占用: 550MB", result.detail)
            self.assertIn("可用: 450MB", result.detail)
            self.assertIn("计算口径: free -/+ buffers/cache", result.detail)

    def test_unparseable_memory_is_critical_instead_of_zero_percent_ok(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(td, "memory.txt", "unsupported free output\n")

            result = _parse_memory(td)[0]
            self.assertEqual(result.status, "CRIT")
            self.assertEqual(result.value, "数据异常")

    def test_sustained_swap_activity_warns_below_usage_threshold(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "memory.txt",
                "              total        used        free      shared  buff/cache   available\n"
                "Mem:          64000       10000        5000           0       49000       40000\n"
                "Swap:         10000        2000        8000\n",
            )
            self._write(
                td, "vmstat.txt",
                " r  b   swpd   free   buff  cache   si   so    bi    bo   in   cs us sy id wa st\n"
                " 0  0 2000000 5000000 1000 10000  0    0     0     0    0    0  1  1 98  0  0\n"
                " 0  0 2000000 5000000 1000 10000 64   32     0     0    0    0  1  1 98  0  0\n"
                " 0  0 2000000 5000000 1000 10000 16    8     0     0    0    0  1  1 98  0  0\n",
            )

            result = next(item for item in _parse_memory(td) if item.name == "Swap使用率")
            self.assertEqual(result.status, "WARN")
            self.assertIn("持续换页: 是", result.detail)
            self.assertIn("检测到持续换页", result.suggestion)

    def test_swap_occupancy_uses_50_and_80_percent_thresholds(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "memory.txt",
                "              total        used        free      shared  buff/cache   available\n"
                "Mem:          64000       10000        5000           0       49000       40000\n"
                "Swap:         10000        5000        5000\n",
            )
            self._write(td, "vmstat.txt", "")

            result = next(item for item in _parse_memory(td) if item.name == "Swap使用率")
            self.assertEqual(result.status, "WARN")
            self.assertEqual(result.value, "50.0%")

            self._write(
                td, "memory.txt",
                "              total        used        free      shared  buff/cache   available\n"
                "Mem:          64000       10000        5000           0       49000       40000\n"
                "Swap:         10000        8000        2000\n",
            )
            result = next(item for item in _parse_memory(td) if item.name == "Swap使用率")
            self.assertEqual(result.status, "CRIT")
            self.assertEqual(result.value, "80.0%")

    def test_iostat_ignores_since_boot_sample(self):
        with tempfile.TemporaryDirectory() as td:
            host = Path(td)
            (host / "meminfo.txt").write_text("HugePages_Total: 1\nHugePages_Free: 1\nHugepagesize: 2048 kB\n", encoding="utf-8")
            (host / "transparent_hugepage.txt").write_text("enabled=[never]\n", encoding="utf-8")
            (host / "process_summary.txt").write_text("S 10\n", encoding="utf-8")
            (host / "iostat.txt").write_text(
                "Device: rrqm/s wrqm/s r/s w/s rkB/s wkB/s avgrq-sz avgqu-sz await r_await w_await svctm %util\n"
                "dm-1 0 0 0 0.1 0 1 10 0 473 0 493 1 0.1\n"
                "Device: rrqm/s wrqm/s r/s w/s rkB/s wkB/s avgrq-sz avgqu-sz await r_await w_await svctm %util\n"
                "sda 0 0 1 1 8 8 16 0.1 5 5 5 1 2\n",
                encoding="utf-8",
            )
            io_result = next(x for x in _parse_oracle_host_readiness(td) if x.name == "磁盘IO延迟")
            self.assertEqual(io_result.status, "OK")
            self.assertIn("5.0ms", io_result.value)

    def test_hugepages_unconfigured_is_ok(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(td, "meminfo.txt", "HugePages_Total: 0\nHugePages_Free: 0\nHugepagesize: 2048 kB\n")
            self._write(td, "transparent_hugepage.txt", "enabled=[never]\n")
            self._write(td, "process_summary.txt", "S 10\n")
            self._write(td, "iostat.txt", "")
            result = next(
                item for item in _parse_oracle_host_readiness(td)
                if item.name == "Oracle HugePages"
            )
            self.assertEqual(result.status, "OK")
            self.assertEqual(result.suggestion, "")
            self.assertIn("未配置", result.detail)

    def test_ntp_unsynced_is_informational(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(td, "datetime.txt", "2026-08-21 17:00:00\n")
            self._write(td, "ntp.txt", "")
            self._write(
                td, "timedatectl.txt",
                "System clock synchronized: no\nNTP service: inactive\n",
            )
            result = _parse_time_sync(td)
            self.assertEqual(result.status, "OK")
            self.assertEqual(result.suggestion, "")
            self.assertEqual(result.value, "未同步")
            self.assertIn("提示", result.detail)

    def test_disabled_firewall_is_acceptable_for_oracle_host(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(td, "firewall.txt", "Active: inactive (dead)\n")
            self._write(td, "iptables.txt", "")
            result = _parse_firewall(td)
            self.assertEqual(result.status, "OK")
            self.assertEqual(result.value, "stopped")
            self.assertEqual(result.suggestion, "")
            self.assertIn("可接受配置", result.detail)

    def test_iptables_accept_policy_is_informational(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(td, "firewall.txt", "Active: inactive (dead)\n")
            self._write(
                td, "iptables.txt",
                "Chain INPUT (policy ACCEPT)\n"
                "target prot opt source destination\n",
            )
            result = _parse_firewall(td)
            self.assertEqual(result.status, "OK")
            self.assertIn("INPUT ACCEPT", result.value)
            self.assertEqual(result.suggestion, "")

    def test_missing_firewall_data_requires_manual_check(self):
        with tempfile.TemporaryDirectory() as td:
            result = _parse_firewall(td)
            self.assertEqual(result.status, "WARN")
            self.assertEqual(result.value, "unknown")
            self.assertIn("人工核对", result.suggestion)

    def test_tablespace_parser_accepts_temp_metrics(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "tablespaces.txt").write_text(
                "TABLESPACE_NAME|CONTENTS|ALLOC_GB|USED_GB|FREE_GB|ALLOC_USAGE_PCT|MAX_USAGE_PCT|MAX_GB|AUT|STATUS\n"
                "USERS|PERMANENT|1|0.2|0.8|20|2|10|YES|ONLINE\n", encoding="utf-8")
            Path(td, "temp_tablespaces.txt").write_text(
                "TABLESPACE_NAME|CONTENTS|ALLOC_GB|USED_GB|FREE_GB|ALLOC_USAGE_PCT|MAX_USAGE_PCT|MAX_GB|AUT|STATUS\n"
                "TEMP|TEMPORARY|1|0|1|0|0|10|YES|ONLINE\n", encoding="utf-8")
            result = _parse_tablespaces(td)[0]
            self.assertEqual(result.status, "OK")
            self.assertIn("2个表空间", result.value)

    def test_autoextensible_tablespace_uses_max_usage_for_status_and_chart(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "tablespaces.txt").write_text(
                "TABLESPACE_NAME|CONTENTS|ALLOC_GB|USED_GB|FREE_GB|ALLOC_USAGE_PCT|MAX_USAGE_PCT|MAX_GB|AUT|STATUS\n"
                "USERS|PERMANENT|99|78|21|99|78|100|YES|ONLINE\n",
                encoding="utf-8",
            )
            result = _parse_tablespaces(td)[0]
            self.assertEqual(result.status, "OK")
            self.assertIn("自动扩展=YES 时按最大使用率判定", result.detail)
            self.assertIn('<div class="bar-value">78.0%</div>', result.extra_html)
            self.assertIn("99.00%", result.extra_html)

    def test_tablespaces_are_sorted_by_max_usage_descending(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "tablespaces.txt",
                "TABLESPACE_NAME|CONTENTS|ALLOC_GB|USED_GB|FREE_GB|ALLOC_USAGE_PCT|MAX_USAGE_PCT|MAX_GB|AUT|STATUS\n"
                "LOW_MAX|PERMANENT|90|80|10|88.9|40|200|YES|ONLINE\n"
                "HIGH_MAX|PERMANENT|20|18|2|90|90|20|NO|ONLINE\n",
            )
            result = _parse_tablespaces(td)[0]
            self.assertLess(result.extra_html.index("HIGH_MAX"), result.extra_html.index("LOW_MAX"))

    def test_force_logging_disabled_is_hint_and_truncated_header_is_ignored(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "db_resilience.txt").write_text(
                "FOR|FLASHBACK_ON|SUPPLEME|DATABASE_ROLE|PROTECTION_MODE|PROTECTION_LEVEL|SWITCHOVER_STATUS\n"
                "NO|YES|NO|PRIMARY|MAXIMUM PERFORMANCE|MAXIMUM PERFORMANCE|NOT ALLOWED\n",
                encoding="utf-8",
            )
            result = _parse_database_resilience(td)[0]
            self.assertEqual(result.status, "OK")
            self.assertEqual(result.value, "存在提示")
            self.assertIn("提示：未启用 FORCE LOGGING", result.detail)
            self.assertEqual(result.suggestion, "")

    def test_actions_are_severity_sorted_and_clickable(self):
        results = {
            "host": [CheckResult("主机警告", "WARN", "1", "")],
            "db": [CheckResult("数据库严重", "CRIT", "2", "")],
        }
        panel = generate_action_panel(results)
        self.assertLess(panel.index("数据库严重"), panel.index("主机警告"))
        self.assertIn(f'href="#{_make_anchor("db", "数据库严重")}"', panel)
        self.assertNotEqual(
            _make_anchor("host", "重复名称"),
            _make_anchor("db", "重复名称"),
        )

    def test_sidebar_contains_every_check_item(self):
        results = {
            "db": [
                CheckResult("实例状态", "OK", "OPEN", ""),
                CheckResult("控制文件", "WARN", "1个", ""),
            ]
        }
        sidebar = generate_sidebar_nav(results)
        self.assertIn(_make_anchor("db", "实例状态"), sidebar)
        self.assertIn(_make_anchor("db", "控制文件"), sidebar)
        self.assertIn('class="dot warn"', sidebar)

    def test_inactive_sessions_keep_total_and_show_top_30(self):
        with tempfile.TemporaryDirectory() as td:
            rows = [
                f"{index}|{index + 100}|USER{index}|INACTIVE|{100-index}|app|host|os|sql{index}"
                for index in range(35)
            ]
            self._write(
                td,
                "dead_processes.txt",
                "SID|SERIAL#|USERNAME|STATUS|INACTIVE_HOURS|PROGRAM|MACHINE|OSUSER|SQL_ID\n"
                + "\n".join(rows),
            )
            result = _parse_dead_processes(td)
            self.assertEqual(result.name, "长时间不活动会话")
            self.assertEqual(result.value, "35个")
            self.assertEqual(result.extra_html.count("<tr>"), 31)
            self.assertIn("仅展示前 30 条", result.extra_html)
            self.assertNotIn("USER34", result.extra_html)

    def test_instance_status_uses_structured_tables(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "instance_status.txt",
                "INSTANCE_NAME|HOST_NAME|STATUS|DATABASE_STATUS|VERSION|STARTUP_TIME\n"
                "orcl|db01|OPEN|ACTIVE|19.0.0.0.0|2026-08-20 10:00:00\n",
            )
            self._write(
                td, "database_status.txt",
                "NAME|OPEN_MODE|DATABASE_ROLE|CREATED|LOG_MODE\n"
                "ORCL|READ WRITE|PRIMARY|2020-01-01|ARCHIVELOG\n",
            )
            result = _parse_instance_status(td)
            self.assertEqual(result.status, "OK")
            self.assertIn("实例信息", result.extra_html)
            self.assertIn("数据库信息", result.extra_html)
            self.assertIn("ARCHIVELOG", result.extra_html)

    def test_instance_status_uses_all_gv_instance_rows(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "instance_status.txt",
                "INST_ID|INSTANCE_NAME|HOST_NAME|STATUS|DATABASE_STATUS|VERSION|STARTUP_TIME\n"
                "1|orcl1|db01|OPEN|ACTIVE|19.0.0.0.0|2026-08-20 10:00:00\n"
                "2|orcl2|db02|OPEN|ACTIVE|19.0.0.0.0|2026-08-20 10:00:01\n",
            )
            result = _parse_instance_status(td)
            self.assertEqual(result.status, "OK")
            self.assertEqual(result.value, "2个实例/全部OPEN")
            self.assertIn("orcl1", result.extra_html)
            self.assertIn("orcl2", result.extra_html)

    def test_data_file_no_autoextend_is_highlighted(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "data_files.txt",
                "FILE_NAME|TABLESPACE_NAME|SIZE_MB|AUTOEXTENSIBLE|MAXSIZE_MB|STATUS|ONLINE_STATUS\n"
                "/oradata/system01.dbf|SYSTEM|1024|NO|1024|AVAILABLE|ONLINE\n",
            )
            result = _parse_data_files(td)
            self.assertEqual(result.status, "WARN")
            self.assertIn('<td class="metric-warn">NO</td>', result.extra_html)

    def test_asm_high_usage_cell_is_highlighted(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "asm_diskgroups.txt",
                "GROUP_NUMBER|NAME|STATE|TYPE|TOTAL_GB|FREE_GB|USAGE_PCT\n"
                "1|DATA|MOUNTED|NORMAL|100|10|90\n",
            )
            result = _parse_asm_diskgroups(td)
            self.assertEqual(result.status, "WARN")
            self.assertIn('<td class="metric-warn">90</td>', result.extra_html)

    def test_control_files_include_full_paths(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "control_files.txt",
                "NAME|STATUS|IS_RECOVERY_DEST_FILE|BLOCK_SIZE|FILE_SIZE_BLKS\n"
                "/u01/oradata/control01.ctl||NO|16384|1024\n"
                "/u02/oradata/control02.ctl||NO|16384|1024\n",
            )
            result = _parse_control_files(td)
            self.assertEqual(result.status, "OK")
            self.assertIn("/u01/oradata/control01.ctl", result.extra_html)
            self.assertIn("16.00", result.extra_html)

    def test_undo_management_includes_extents_and_parameters(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "undo_tablespace.txt",
                "TABLESPACE_NAME|STATUS|SIZE_MB\nUNDOTBS1|ACTIVE|128\nUNDOTBS1|EXPIRED|256\n",
            )
            self._write(
                td, "undo_retention.txt",
                "NAME|VALUE\nundo_retention|900\nundo_tablespace|UNDOTBS1\n",
            )
            result = _parse_rollback_segments(td)
            self.assertEqual(result.name, "UNDO/回滚段管理")
            self.assertIn("ACTIVE", result.extra_html)
            self.assertIn("undo_retention", result.extra_html)

    def test_hit_ratios_include_health_graphics(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "buffer_cache_hit.txt",
                "NAME|PHYSICAL_READS|DB_BLOCK_GETS|CONSISTENT_GETS|HIT_RATIO\n"
                "DEFAULT|100|1000|9000|99\n",
            )
            self._write(
                td, "library_cache_hit.txt",
                "PINS|RELOADS|HIT_RATIO\n10000|10|99.9\n",
            )
            buffer_result = _parse_buffer_cache_hit(td)
            library_result = _parse_library_cache_hit(td)
            self.assertIn("health-percentage", buffer_result.extra_html)
            self.assertIn("bar-fill-ok", buffer_result.extra_html)
            self.assertIn("DEFAULT", buffer_result.extra_html)
            self.assertIn("health-percentage", library_result.extra_html)
            self.assertIn("Reloads", library_result.extra_html)

    def test_buffer_cache_numbers_do_not_use_scientific_notation(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "buffer_cache_hit.txt",
                "NAME|PHYSICAL_READS|DB_BLOCK_GETS|CONSISTENT_GETS|HIT_RATIO\n"
                "DEFAULT|1.23E+09|2.00E+10|3.40E+10|97.81\n",
            )
            result = _parse_buffer_cache_hit(td)
            self.assertNotIn("E+", result.extra_html)
            self.assertIn("1230000000", result.extra_html)

    def test_inspection_explanation_is_shared_by_html(self):
        section = generate_category_section(
            "db", [CheckResult(
                "SYSTEM表空间业务对象", "WARN", "1个", "发现对象",
                extra_html='<table class="data-table"><tr><td>明细对象</td></tr></table>',
            )]
        )
        self.assertIn("业务对象误放", section)
        self.assertIn("巡检说明", section)
        self.assertLess(section.index("明细对象"), section.index("巡检说明"))
        self.assertIn(
            '<div class="extra-content"><table class="data-table">', section
        )

    def test_explanation_is_collapsed_even_without_other_details(self):
        section = generate_category_section(
            "db", [CheckResult("实例状态", "OK", "OPEN", "状态正常")]
        )
        self.assertIn("&#9656; 展开详情", section)
        self.assertIn('<div class="extra-content"><div class="check-purpose">', section)
        self.assertLess(section.index("检查结果"), section.index("巡检说明"))

    def test_top_disk_read_sql_includes_sql_text(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "top_disk_read_sql.txt",
                "SQL_ID|DISK_READS|EXECUTIONS|READS_PER_EXEC|BUFFER_GETS|PARSING_SCHEMA_NAME|SQL_TEXT\n"
                "abc123|50000|10|5000|90000|APP|select * from orders\n",
            )
            result = _parse_top_disk_read_sql(td)
            self.assertIn("abc123", result.extra_html)
            self.assertIn("select * from orders", result.extra_html)

    def test_sql92_false_is_informational(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "security_parameters.txt",
                "NAME|VALUE|ISDEFAULT\n"
                "sql92_security|FALSE|TRUE\n"
                "remote_os_authent|FALSE|TRUE\n",
            )
            result = next(
                item for item in _parse_audit_and_access_controls(td)
                if item.name == "安全参数"
            )
            self.assertEqual(result.status, "OK")
            self.assertIn("sql92_security=FALSE", result.detail)

    def test_listener_optional_hardening_is_informational(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(
                td, "listener_status.txt",
                "status READY\nSecurity ON: Local OS Authentication\n"
                "The command completed successfully\n",
            )
            self._write(td, "listener_ora.txt", "LISTENER=(DESCRIPTION=(ADDRESS=(PROTOCOL=TCP)))\n")
            self._write(td, "sqlnet_ora.txt", "SQLNET.ENCRYPTION_SERVER=ACCEPTED\n")
            result = _parse_listener_security(td)
            self.assertEqual(result.status, "OK")
            self.assertEqual(result.suggestion, "")
            self.assertIn("ADMIN_RESTRICTIONS=未配置", result.detail)
            self.assertIn("传输加密=未要求", result.detail)

    def test_listener_failure_remains_critical(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(td, "listener_status.txt", "TNS-12541: no listener\n")
            result = _parse_listener_security(td)
            self.assertEqual(result.status, "CRIT")
            self.assertIn("未处于 READY", result.suggestion)

    def test_html_cells_are_escaped(self):
        table = generate_data_table(["<列>"], [["<script>alert(1)</script>"]])
        self.assertNotIn("<script>", table)
        self.assertIn("&lt;script&gt;", table)

    def test_tar_path_traversal_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            archive = Path(td, "bad.tar.gz")
            with tarfile.open(archive, "w:gz") as tar:
                info = tarfile.TarInfo("../escape.txt")
                payload = b"bad"
                info.size = len(payload)
                tar.addfile(info, io.BytesIO(payload))
            with self.assertRaises(ValueError):
                extract_tar_gz(str(archive))

    def test_archive_temp_directory_is_cleaned_after_report(self):
        with tempfile.TemporaryDirectory() as td:
            raw = Path(self._make_collection(td, "raw", "app01", kind="host"))
            archive = Path(td, "host_check_app01.tar.gz")
            with tarfile.open(archive, "w:gz") as tar:
                for path in raw.rglob("*"):
                    tar.add(path, arcname=path.relative_to(raw))
            temp_root = Path(tempfile.gettempdir())
            before = {path.resolve() for path in temp_root.glob("oracle_check_*")}
            output = Path(td, "report.html")
            build_report([str(archive)], str(output), write_docx=False)
            after = {path.resolve() for path in temp_root.glob("oracle_check_*")}
            self.assertEqual(after, before)

    def test_template_path_exists(self):
        self.assertTrue(os.path.isfile(template_path()), template_path())

    def test_source_mode_uses_reporter_as_project_root(self):
        self.assertEqual(Path(project_root()).resolve(), ROOT.resolve())

    def test_report_name_uses_sid_and_date(self):
        env = {"oracle_sid": "orcl", "timestamp": "20260821_101500", "hostname": "db01"}
        self.assertEqual(report_basename(env), "Oracle巡检报告_orcl_20260821.html")
        self.assertEqual(
            report_basename({"hostname": "db01", "timestamp": "2026-08-21 10:15:00"}),
            "主机巡检报告_db01_20260821.html",
        )
        self.assertEqual(
            report_basename({"timestamp": "20260821_101500"}),
            "主机巡检报告_unknown_20260821.html",
        )
        with tempfile.TemporaryDirectory() as td:
            named = default_output_path(env, td)
            self.assertEqual(Path(named).name, "Oracle巡检报告_orcl_20260821.html")
            host_named = default_output_path(
                {"hostname": "app01", "timestamp": "20260821_101500"}, td
            )
            self.assertEqual(Path(host_named).name, "主机巡检报告_app01_20260821.html")
            rewritten = resolve_output_file(env, str(Path(td, "oracle_inspection.html")))
            self.assertEqual(Path(rewritten).name, "Oracle巡检报告_orcl_20260821.html")
            placeholder = resolve_output_file(env, str(Path(td, "Oracle巡检报告.html")))
            self.assertEqual(Path(placeholder).name, "Oracle巡检报告_orcl_20260821.html")
            host_placeholder = resolve_output_file(
                {"hostname": "app01", "timestamp": "20260821_101500"},
                str(Path(td, "主机巡检报告.html")),
            )
            self.assertEqual(Path(host_placeholder).name, "主机巡检报告_app01_20260821.html")
            previous = resolve_output_file(
                {"oracle_sid": "prod", "timestamp": "20260822_090000"},
                str(Path(td, "Oracle巡检报告_orcl_20260821.html")),
            )
            self.assertEqual(Path(previous).name, "Oracle巡检报告_prod_20260822.html")
            previous_host = resolve_output_file(
                {"hostname": "app02", "timestamp": "20260822_090000"},
                str(Path(td, "主机巡检报告_app01_20260821.html")),
            )
            self.assertEqual(Path(previous_host).name, "主机巡检报告_app02_20260822.html")
            custom = resolve_output_file(env, str(Path(td, "custom.html")))
            self.assertEqual(Path(custom).name, "custom.html")
            as_dir = resolve_output_file(env, td)
            self.assertEqual(Path(as_dir), Path(td, "Oracle巡检报告_orcl_20260821.html"))

    def test_gui_path_helpers_and_score_tones(self):
        with tempfile.TemporaryDirectory() as td:
            first = str(Path(td, "host_check.tar.gz"))
            Path(first).write_text("", encoding="utf-8")
            paths, added = add_unique_paths([], [first, first])
            self.assertEqual(added, 1)
            self.assertEqual(paths, [os.path.normpath(first)])
            self.assertEqual(classify_input(td), "数据目录")
            self.assertEqual(classify_input(first), "采集包")
            self.assertEqual(
                normalize_output_selection(td),
                os.path.normpath(td),
            )
            self.assertEqual(
                normalize_output_selection(str(Path(td, "report.html"))),
                os.path.normpath(td),
            )
            self.assertEqual(
                normalize_output_selection(str(Path(td, "report.docx"))),
                os.path.normpath(td),
            )
            self.assertEqual(score_tone(80), "success")
            self.assertEqual(score_tone(79), "warning")
            self.assertEqual(score_tone(60), "warning")
            self.assertEqual(score_tone(59), "danger")
            self.assertEqual(score_tone(96), "success")
            self.assertEqual(score_tone(96, has_score=False), "muted")
            self.assertEqual(score_band(80), "OK")
            self.assertEqual(score_band(79), "WARN")
            self.assertEqual(score_band(60), "WARN")
            self.assertEqual(score_band(59), "CRIT")

    def test_drag_drop_accepts_archives_and_folders_only(self):
        with tempfile.TemporaryDirectory() as td:
            folder = Path(td, "采集数据")
            folder.mkdir()
            archive = Path(td, "db check orcl.tar.gz")
            archive.write_bytes(b"test")
            tgz = Path(td, "host_check.tgz")
            tgz.write_bytes(b"test")
            text_file = Path(td, "notes.txt")
            text_file.write_text("ignore", encoding="utf-8")
            missing = Path(td, "missing.tar.gz")
            accepted, rejected = filter_dropped_paths(
                [str(folder), str(archive), str(tgz), str(text_file), str(missing)]
            )
            self.assertEqual(
                accepted,
                [os.path.normpath(str(folder)), os.path.normpath(str(archive)),
                 os.path.normpath(str(tgz))],
            )
            self.assertEqual(
                rejected,
                [os.path.normpath(str(text_file)), os.path.normpath(str(missing))],
            )

    def test_choose_window_geometry_fits_desktop_and_laptop(self):
        width, height, x, y = choose_window_geometry(900, 700, 1920, 1080)
        self.assertEqual((width, height), (1180, 820))
        self.assertEqual(x, (1920 - 1180) // 2)
        self.assertGreaterEqual(y, 0)

        width, height, x, y = choose_window_geometry(1200, 900, 1366, 768)
        self.assertEqual(width, 1200)
        self.assertEqual(height, 768 - 88)
        self.assertGreaterEqual(x, 0)
        self.assertGreaterEqual(y, 0)

        width, height, x, y = choose_window_geometry(900, 700, 1280, 720)
        self.assertEqual(width, 1180)
        self.assertLessEqual(height, 720 - 88)
        self.assertGreaterEqual(height, 520)

    def test_build_report_requires_input(self):
        with self.assertRaises(ReportBuildError):
            build_report([])

    def test_build_report_rejects_empty_collection(self):
        with tempfile.TemporaryDirectory() as td:
            Path(td, "env.info").write_text("hostname=test\n", encoding="utf-8")
            with self.assertRaises(ReportBuildError):
                build_report([td])

    def test_build_report_respects_selected_formats(self):
        with tempfile.TemporaryDirectory() as td:
            raw = Path(td, "raw")
            raw.mkdir()
            (raw / "host").mkdir()
            self._write(raw, "env.info", "hostname=fmt-smoke\ntimestamp=20260821_101500\n")
            self._write(
                raw,
                "collection_manifest.tsv",
                "item\ttype\tstatus\texit_code\tmessage\n"
                "missing.txt\tCOMMAND\tFAILED\t1\tsmoke\n",
            )
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            with self.assertRaises(ReportBuildError):
                build_report([str(raw)], str(out_dir), write_html=False, write_docx=False)

            html_only = build_report(
                [str(raw)], str(out_dir), write_html=True, write_docx=False
            )
            self.assertTrue(html_only.html and Path(html_only.html).is_file())
            self.assertIsNone(html_only.docx)
            self.assertFalse(Path(html_only.html).with_suffix(".docx").exists())

            docx_only = build_report(
                [str(raw)], str(out_dir), write_html=False, write_docx=True
            )
            self.assertIsNone(docx_only.html)
            self.assertTrue(docx_only.docx and Path(docx_only.docx).is_file())
            self.assertEqual(Path(docx_only.docx).parent, out_dir)

    def _make_collection(self, root, name, hostname, *, sid=None, kind="host",
                         timestamp="20260821_101500", schema_version=None):
        raw = Path(root, name)
        raw.mkdir()
        env = f"hostname={hostname}\ntimestamp={timestamp}\n"
        if schema_version:
            env += f"schema_version={schema_version}\n"
        if kind == "host":
            (raw / "host").mkdir()
        else:
            (raw / "db").mkdir()
            (raw / "security").mkdir()
            env += f"oracle_sid={sid}\n"
        self._write(raw, "env.info", env)
        self._write(
            raw, "collection_manifest.tsv",
            "item\ttype\tstatus\texit_code\tmessage\n",
        )
        return str(raw)

    def test_expand_input_paths_unfolds_archive_folder(self):
        with tempfile.TemporaryDirectory() as td:
            first = Path(td, "host_check_a.tar.gz")
            second = Path(td, "db_check_orcl.tar.gz")
            first.write_bytes(b"not-a-real-tar")
            second.write_bytes(b"not-a-real-tar")
            expanded = expand_input_paths([td])
            self.assertEqual(
                [Path(path).name for path in expanded],
                ["db_check_orcl.tar.gz", "host_check_a.tar.gz"],
            )

    def test_group_report_jobs_keeps_host_and_database_separate(self):
        with tempfile.TemporaryDirectory() as td:
            host = self._make_collection(td, "host1", "app01", kind="host")
            db_orcl = self._make_collection(
                td, "db_orcl", "app01", sid="orcl", kind="db"
            )
            db_prod = self._make_collection(
                td, "db_prod", "app01", sid="prod", kind="db"
            )
            separate = group_report_jobs([host, db_orcl])
            self.assertEqual(len(separate), 2)
            self.assertTrue(all(len(job) == 1 for job in separate))
            self.assertEqual([job[0].report_kind for job in separate], ["host", "db"])
            jobs = group_report_jobs([host, db_orcl, db_prod])
            self.assertEqual(len(jobs), 3)
            sids = {job[0].env.get("oracle_sid") for job in jobs if job[0].has_db}
            self.assertEqual(sids, {"orcl", "prod"})
            self.assertEqual(sum(job[0].has_host for job in jobs), 1)

            two_dbs = group_report_jobs([db_orcl, db_prod])
            self.assertEqual(len(two_dbs), 2)

    def test_group_report_jobs_never_pairs_different_hosts(self):
        with tempfile.TemporaryDirectory() as td:
            host = self._make_collection(td, "host1", "app01", kind="host")
            db = self._make_collection(td, "db1", "app02", sid="orcl", kind="db")
            jobs = group_report_jobs([host, db])
            self.assertEqual(len(jobs), 2)
            self.assertTrue(all(len(job) == 1 for job in jobs))

    def test_group_report_jobs_never_reuses_host_snapshot_for_database(self):
        with tempfile.TemporaryDirectory() as td:
            old_host = self._make_collection(
                td, "host_old", "app01", kind="host", timestamp="20260820_080000"
            )
            close_host = self._make_collection(
                td, "host_close", "app01", kind="host", timestamp="20260821_100000"
            )
            db = self._make_collection(
                td, "db", "app01", sid="orcl", kind="db", timestamp="20260821_101500"
            )
            jobs = group_report_jobs([old_host, close_host, db])
            self.assertEqual(len(jobs), 3)
            self.assertTrue(all(len(job) == 1 for job in jobs))
            self.assertEqual([job[0].source_path for job in jobs], [old_host, close_host, db])
            self.assertEqual([job[0].report_kind for job in jobs], ["host", "host", "db"])

    def test_combined_directory_is_split_into_host_and_database_jobs(self):
        with tempfile.TemporaryDirectory() as td:
            raw = Path(td, "combined")
            raw.mkdir()
            (raw / "host").mkdir()
            (raw / "db").mkdir()
            (raw / "security").mkdir()
            self._write(
                raw, "env.info",
                "hostname=app01\noracle_sid=orcl\ntimestamp=20260821_101500\n",
            )
            self._write(
                raw, "collection_manifest.tsv",
                "item\ttype\tstatus\texit_code\tmessage\n",
            )
            jobs = group_report_jobs([str(raw)])
            self.assertEqual(len(jobs), 2)
            self.assertEqual([job[0].report_kind for job in jobs], ["host", "db"])
            self.assertTrue(jobs[0][0].has_host and not jobs[0][0].has_db)
            self.assertTrue(jobs[1][0].has_db and not jobs[1][0].has_host)

    def test_combined_directory_integrity_failures_are_scoped_per_report(self):
        with tempfile.TemporaryDirectory() as td:
            raw = Path(td, "combined")
            raw.mkdir()
            for section in ("host", "db", "security"):
                (raw / section).mkdir()
            self._write(
                raw, "env.info",
                "hostname=app01\noracle_sid=orcl\ntimestamp=20260821_101500\n",
            )
            self._write(raw / "host", "host_failed.txt", "host command failed\n")
            self._write(raw / "db", "db_failed.txt", "database command failed\n")
            self._write(
                raw, "collection_manifest.tsv",
                "item\ttype\tstatus\texit_code\tmessage\n"
                "host_failed.txt\tCOMMAND\tFAILED\t1\thost failure\n"
                "db_failed.txt\tSQL\tFAILED\t1\tdatabase failure\n",
            )
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            batch = build_reports(
                [str(raw)], str(out_dir), write_docx=False, write_summary=False
            )
            self.assertEqual(len(batch.reports), 2)
            host_report = next(report for report in batch.reports if "host" in report.results)
            db_report = next(report for report in batch.reports if "db" in report.results)
            host_integrity = host_report.results["host"][0]
            db_integrity = db_report.results["db"][0]
            self.assertEqual(host_integrity.value, "1项失败")
            self.assertIn("host_failed.txt", host_integrity.extra_html)
            self.assertNotIn("db_failed.txt", host_integrity.extra_html)
            self.assertEqual(db_integrity.value, "1项失败")
            self.assertIn("db_failed.txt", db_integrity.extra_html)
            self.assertNotIn("host_failed.txt", db_integrity.extra_html)

    def test_future_schema_version_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            raw = self._make_collection(
                td, "future", "app01", kind="host", schema_version="5.0"
            )
            with self.assertRaisesRegex(ValueError, "不支持采集包协议"):
                group_report_jobs([raw])

    def test_health_score_respects_severity(self):
        self.assertEqual(
            calculate_health_score({"OK": 100, "WARN": 0, "CRIT": 0, "TOTAL": 100}),
            100,
        )
        self.assertEqual(
            calculate_health_score({"OK": 9, "WARN": 1, "CRIT": 0, "TOTAL": 10}),
            96,
        )
        self.assertEqual(
            calculate_health_score({"OK": 9, "WARN": 0, "CRIT": 1, "TOTAL": 10}),
            90,
        )

        weighted = {
            "db": [
                CheckResult("数据库版本", "OK", "", ""),
                CheckResult("表空间使用概览", "WARN", "", ""),
                CheckResult("无效对象", "CRIT", "", ""),
                CheckResult("AWR分析", "INFO", "未启用", ""),
            ]
        }
        breakdown = calculate_score_breakdown(weighted)
        self.assertEqual(breakdown.scored_items, 3)
        self.assertLess(breakdown.score, 100)
        self.assertEqual(breakdown.confidence, "无法验证")

    def test_high_score_color_is_independent_from_critical_risk(self):
        results = {
            "db": [
                CheckResult("采集完整性", "INFO", "全部成功", ""),
                *[
                    CheckResult(f"常规检查{i}", "OK", "正常", "")
                    for i in range(9)
                ],
                CheckResult("严重检查", "CRIT", "异常", "立即处理"),
            ]
        }
        counts = {"OK": 9, "WARN": 0, "CRIT": 1, "INFO": 1, "UNKNOWN": 0, "TOTAL": 11}
        breakdown = calculate_score_breakdown(results)
        self.assertEqual(breakdown.score, 90)

        with tempfile.TemporaryDirectory() as td:
            html_path = Path(td, "report.html")
            docx_path = Path(td, "report.docx")
            generate_report(
                results,
                {"hostname": "db01", "oracle_sid": "orcl", "timestamp": "20260827_120000"},
                template_path(),
                str(html_path),
            )
            generate_docx_report(
                results,
                {"hostname": "db01", "oracle_sid": "orcl", "timestamp": "20260827_120000"},
                str(docx_path),
                breakdown.score,
                counts,
            )

            report = html_path.read_text(encoding="utf-8")
            self.assertIn('<div class="score-ring good">90</div>', report)
            self.assertIn('class="badge badge-crit"', report)
            self.assertIn("存在严重问题", report)

            document = Document(docx_path)
            score_paragraph = next(
                paragraph for paragraph in document.paragraphs
                if "健康评分 / 100" in paragraph.text
            )
            score_run = next(run for run in score_paragraph.runs if run.text == "90")
            risk_run = next(
                run for run in score_paragraph.runs if run.text == "存在严重问题"
            )
            self.assertEqual(str(score_run.font.color.rgb), DOCX_COLORS["OK"])
            self.assertEqual(str(risk_run.font.color.rgb), DOCX_COLORS["CRIT"])

        incomplete = {
            "db": [
                CheckResult("采集完整性", "UNKNOWN", "1项失败", ""),
                CheckResult("数据库版本", "OK", "", ""),
            ]
        }
        incomplete_breakdown = calculate_score_breakdown(incomplete)
        self.assertEqual(incomplete_breakdown.score, 100)
        self.assertEqual(incomplete_breakdown.confidence, "不完整")

    def test_awr_config_skip_is_informational(self):
        with tempfile.TemporaryDirectory() as td:
            self._write(td, "awr_snapshot.txt", "SKIPPED|按配置未启用 AWR 采集\n")
            result = _parse_awr_metrics(td)[0]
            self.assertEqual(result.status, "INFO")
            self.assertEqual(result.value, "未启用")

    def test_build_reports_writes_one_report_per_collection_group(self):
        with tempfile.TemporaryDirectory() as td:
            first = self._make_collection(td, "raw_a", "app01", sid="orcl", kind="db")
            second = self._make_collection(td, "raw_b", "app02", sid="prod", kind="db")
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            batch = build_reports([first, second], str(out_dir), write_docx=False)
            self.assertEqual(len(batch.reports), 2)
            self.assertEqual(batch.errors, [])
            html_files = sorted(path.name for path in out_dir.glob("*.html"))
            self.assertEqual(
                html_files,
                [
                    "Oracle巡检报告_orcl_20260821.html",
                    "Oracle巡检报告_prod_20260821.html",
                    "Oracle巡检汇总摘要_20260821.html",
                ],
            )
            self.assertTrue(batch.summary_html)
            self.assertIsNone(batch.summary_docx)
            self.assertEqual(list(out_dir.glob("*.docx")), [])

    def test_build_reports_can_stop_before_first_report(self):
        with tempfile.TemporaryDirectory() as td:
            raw = self._make_collection(td, "raw", "app01", kind="host")
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            batch = build_reports(
                [raw], str(out_dir), write_docx=False,
                cancel_check=lambda: True,
            )
            self.assertTrue(batch.cancelled)
            self.assertEqual(batch.reports, [])
            self.assertEqual(list(out_dir.glob("*.html")), [])

    def test_build_reports_stop_preserves_completed_reports(self):
        with tempfile.TemporaryDirectory() as td:
            first = self._make_collection(td, "raw_a", "app01", kind="host")
            second = self._make_collection(td, "raw_b", "app02", kind="host")
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            state = {"cancel": False}

            def progress(index, _total, _label):
                if index == 2:
                    state["cancel"] = True

            batch = build_reports(
                [first, second], str(out_dir), write_docx=False,
                write_summary=False, progress=progress,
                cancel_check=lambda: state["cancel"],
            )
            self.assertTrue(batch.cancelled)
            self.assertEqual(len(batch.reports), 1)
            self.assertTrue(Path(batch.reports[0].html).is_file())
            self.assertEqual(len(list(out_dir.glob("*.html"))), 1)
            self.assertIsNone(batch.summary_html)

    def test_host_and_database_reports_have_isolated_categories(self):
        with tempfile.TemporaryDirectory() as td:
            host = self._make_collection(td, "host", "app01", kind="host")
            db = self._make_collection(td, "db", "app01", sid="orcl", kind="db")
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            batch = build_reports([host, db], str(out_dir), write_docx=False)
            self.assertEqual(len(batch.reports), 2)
            self.assertIsNone(batch.summary_html)
            host_report = next(report for report in batch.reports if "host" in report.results)
            db_report = next(report for report in batch.reports if "db" in report.results)
            self.assertEqual(set(host_report.results), {"host"})
            self.assertNotIn("host", db_report.results)
            self.assertIn("security", db_report.results)
            self.assertEqual(Path(host_report.html).name, "主机巡检报告_app01_20260821.html")
            self.assertEqual(Path(db_report.html).name, "Oracle巡检报告_orcl_20260821.html")
            self.assertNotIn('id="section-host"', Path(db_report.html).read_text(encoding="utf-8"))
            self.assertNotIn('id="section-db"', Path(host_report.html).read_text(encoding="utf-8"))

    def test_docx_summary_chapter_is_fourth_when_three_categories_exist(self):
        results = {
            "host": [CheckResult("CPU", "OK", "10%", "正常")],
            "db": [CheckResult("实例状态", "OK", "OPEN", "正常")],
            "security": [CheckResult("监听器", "OK", "READY", "正常")],
        }
        counts = {"OK": 3, "WARN": 0, "CRIT": 0, "TOTAL": 3}
        with tempfile.TemporaryDirectory() as td:
            output = Path(td, "summary.docx")
            generate_docx_report(results, {"hostname": "db01"}, str(output), 100, counts)
            document = Document(output)
            heads = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style.name == "Heading 1"
            ]
            self.assertIn("1. 主机巡检", heads)
            self.assertIn("2. 数据库巡检", heads)
            self.assertIn("3. 安全巡检", heads)
            self.assertIn("4. 总结", heads)
            self.assertEqual(heads[-1], "4. 总结")
            summary_index = next(
                index for index, paragraph in enumerate(document.paragraphs)
                if paragraph.text == "4. 总结"
            )
            self.assertFalse(any(
                paragraph.text.strip()
                for paragraph in document.paragraphs[summary_index + 1:]
            ))

    def test_optional_report_summary_fills_html_and_word_only_when_enabled(self):
        results = {
            "db": [
                CheckResult("实例状态", "OK", "OPEN", "实例运行正常"),
                CheckResult(
                    "表空间使用率", "CRIT", "MAX 使用率 96%", "表空间使用率过高",
                    "尽快扩容并复查自动扩展配置",
                ),
                CheckResult(
                    "Redo Log", "WARN", "切换 12 次/小时", "切换频率偏高",
                    "适当增加 Redo Log 大小",
                ),
            ]
        }
        counts = {"OK": 1, "WARN": 1, "CRIT": 1, "TOTAL": 3}
        with tempfile.TemporaryDirectory() as td:
            blank_docx = Path(td, "blank.docx")
            filled_docx = Path(td, "filled.docx")
            generate_docx_report(
                results, {"hostname": "db01"}, str(blank_docx), 20, counts
            )
            generate_docx_report(
                results, {"hostname": "db01"}, str(filled_docx), 20, counts,
                generate_summary_content=True,
            )
            blank_paragraphs = Document(blank_docx).paragraphs
            filled_paragraphs = Document(filled_docx).paragraphs
            blank_summary_index = max(
                index for index, paragraph in enumerate(blank_paragraphs)
                if paragraph.text.endswith(". 总结")
            )
            blank_summary_text = "\n".join(
                paragraph.text for paragraph in blank_paragraphs[blank_summary_index + 1:]
            )
            filled_text = "\n".join(p.text for p in filled_paragraphs)
            self.assertNotIn("总体结论", blank_summary_text)
            self.assertNotIn("重点问题", blank_summary_text)
            self.assertIn("总体结论", filled_text)
            self.assertIn("重点问题", filled_text)
            self.assertIn("[严重] 数据库巡检 / 表空间使用率", filled_text)
            self.assertIn("尽快扩容并复查自动扩展配置", filled_text)

            blank_html = Path(td, "blank.html")
            filled_html = Path(td, "filled.html")
            env = {"hostname": "db01", "timestamp": "20260825_120000"}
            generate_report(results, env, template_path(), str(blank_html))
            generate_report(
                results, env, template_path(), str(filled_html),
                generate_summary_content=True,
            )
            blank_report = blank_html.read_text(encoding="utf-8")
            filled_report = filled_html.read_text(encoding="utf-8")
            self.assertIn('id="report-summary"', blank_report)
            self.assertNotIn('<div class="report-summary-content">', blank_report)
            self.assertIn('<div class="report-summary-content">', filled_report)
            self.assertIn("MAX 使用率 96%", filled_report)

    def test_generated_report_contains_navigation_behavior(self):
        with tempfile.TemporaryDirectory() as td:
            raw = Path(td, "raw")
            raw.mkdir()
            (raw / "host").mkdir()
            self._write(raw, "env.info", "hostname=ui-smoke\ntimestamp=20260820_154500\n")
            self._write(
                raw,
                "collection_manifest.tsv",
                "item\ttype\tstatus\texit_code\tmessage\n"
                "missing.txt\tCOMMAND\tFAILED\t1\tsmoke\n",
            )
            output = Path(td, "html.reports", "report.HTML")
            output.parent.mkdir(parents=True)
            (output.parent / ".gitkeep").write_text("", encoding="utf-8")
            output.with_suffix(".json").write_text("{}", encoding="utf-8")
            outputs = build_report([str(raw)], str(output))
            report = Path(outputs.html).read_text(encoding="utf-8")
            self.assertIn('class="action-item" href="#check-host-', report)
            self.assertIn('class="nav-link nav-check"', report)
            self.assertNotIn("${OVERVIEW_NAV}", report)
            self.assertNotIn("${DETAIL_NAV_LABEL}", report)
            self.assertIn("巡检详情", report)
            self.assertIn("IntersectionObserver", report)
            self.assertIn(".nav-link.active", report)
            self.assertTrue(Path(outputs.html).is_file())
            self.assertTrue(Path(outputs.docx).is_file())
            self.assertTrue(Path(outputs.html).with_suffix(".json").exists())
            self.assertEqual(Path(outputs.html).suffix, ".html")
            self.assertEqual(Path(outputs.docx).parent, Path(outputs.html).parent)
            self.assertTrue((Path(outputs.html).parent / ".gitkeep").exists())

            document = Document(outputs.docx)
            item_headings = [
                re.sub(r"^\d+\.\d+\s+", "", paragraph.text)
                for paragraph in document.paragraphs
                if paragraph.style.name == "Heading 2"
                and re.match(r"^\d+\.\d+\s+", paragraph.text)
            ]
            result_items = [
                item.name
                for category_items in outputs.results.values()
                for item in category_items
            ]
            self.assertEqual(len(item_headings), outputs.counts["TOTAL"])
            self.assertCountEqual(item_headings, result_items)
            self.assertEqual(document.paragraphs[0].text, "")
            self.assertNotIn('w:type="page"', document.paragraphs[0]._p.xml)
            self.assertEqual(
                next(
                    paragraph.text for paragraph in document.paragraphs
                    if paragraph.style.name == "Title"
                ),
                document.core_properties.title,
            )
            self.assertNotIn("w:updateFields", document.settings._element.xml)
            self.assertNotIn("TOC", document.element.body.xml)
            self.assertIn("PAGEREF report_category_1", document.element.body.xml)
            self.assertNotIn("章节索引", [paragraph.text for paragraph in document.paragraphs])
            self.assertIn("目录", [paragraph.text for paragraph in document.paragraphs])
            self.assertIn('w:anchor="report_category_1"', document.element.body.xml)
            self.assertIn('w:name="report_category_1"', document.element.body.xml)
            paragraphs = document.paragraphs
            summary_position, summary_heading = next(
                (index, paragraph) for index, paragraph in enumerate(paragraphs)
                if paragraph.style.name == "Heading 1"
                and re.match(r"^\d+\. 总结$", paragraph.text)
            )
            self.assertTrue(summary_heading.paragraph_format.page_break_before)
            self.assertNotIn(
                'w:type="page"',
                paragraphs[summary_position - 1]._p.xml,
            )
            self.assertIn("PAGE", "".join(
                footer._element.xml for footer in (section.footer for section in document.sections)
            ))
            with zipfile.ZipFile(outputs.docx) as package:
                external_relationships = [
                    name for name in package.namelist()
                    if name.endswith(".rels")
                    and 'TargetMode="External"' in package.read(name).decode("utf-8", errors="ignore")
                ]
            self.assertEqual(external_relationships, [])

    def test_extra_html_maps_to_word_blocks_without_omission(self):
        rich_html = (
            '<div class="detail-subtitle">容量明细</div>'
            '<table class="data-table"><thead><tr><th>名称</th><th>值</th></tr></thead>'
            '<tbody><tr><td>USERS</td><td>80%</td></tr></tbody></table>'
            '<div class="bar-chart"><div class="bar-row">'
            '<div class="bar-label">USERS</div><div class="bar-track">'
            '<div class="bar-fill bar-fill-warn" style="width:80%"></div></div>'
            '<div class="bar-value">80.0%</div></div></div>'
            '<ul class="detail-list"><li>第一项</li><li>第二项</li></ul>'
            '<div class="detail-note">仅展示前 30 条</div>'
        )
        blocks = parse_extra_html(rich_html)
        kinds = [block[0] for block in blocks]
        self.assertEqual(kinds, ["subtitle", "table", "bar", "bullets", "note"])
        self.assertIn("USERS", str(blocks))
        self.assertIn("第一项", str(blocks))

    def test_docx_renders_rich_detail_for_every_check(self):
        extra_html = (
            '<div class="detail-subtitle">SQL明细</div>'
            '<table class="data-table"><thead><tr><th>SQL_ID</th><th>SQL文本</th></tr></thead>'
            '<tbody><tr><td>abc123</td><td>select * from orders</td></tr></tbody></table>'
            '<div class="health-percentage"><div class="health-percentage-head">'
            '<span>命中率</span><strong>99.0%</strong></div>'
            '<div class="health-percentage-fill bar-fill-ok"></div></div>'
        )
        results = {
            "db": [
                CheckResult("实例状态", "OK", "OPEN", "实例运行正常"),
                CheckResult(
                    "Disk Read最高SQL", "WARN", "1条", "存在高读取SQL",
                    "建议优化SQL", extra_html=extra_html,
                ),
            ]
        }
        counts = {"OK": 1, "WARN": 1, "CRIT": 0, "TOTAL": 2}
        env = {
            "hostname": "db01",
            "server_ip": "10.0.0.1",
            "oracle_sid": "ORCL",
            "timestamp": "20260820_160000",
        }
        with tempfile.TemporaryDirectory() as td:
            output = Path(td, "enterprise.docx")
            rendered = generate_docx_report(results, env, str(output), 50, counts)
            self.assertEqual(rendered, counts["TOTAL"])
            document = Document(output)
            all_text = "\n".join(
                paragraph.text for paragraph in document.paragraphs
            ) + "\n" + "\n".join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("1.1  实例状态", all_text)
            self.assertIn("1.2  Disk Read最高SQL", all_text)
            self.assertIn("2. 总结", all_text)
            summary_heads = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style.name == "Heading 1" and paragraph.text.endswith("总结")
            ]
            self.assertEqual(summary_heads, ["2. 总结"])
            self.assertIn("abc123", all_text)
            self.assertIn("select * from orders", all_text)
            self.assertIn("建议优化SQL", all_text)

    def test_docx_uses_simsun_for_chinese_and_times_for_latin(self):
        results = {
            "db": [CheckResult("实例状态 Instance 1", "OK", "OPEN 19c", "运行正常")]
        }
        counts = {"OK": 1, "WARN": 0, "CRIT": 0, "TOTAL": 1}
        with tempfile.TemporaryDirectory() as td:
            output = Path(td, "fonts.docx")
            generate_docx_report(
                results,
                {
                    "hostname": "db01",
                    "oracle_sid": "ORCL",
                    "timestamp": "20260821_120000",
                },
                str(output),
                100,
                counts,
            )
            document = Document(output)

            def assert_bilingual_fonts(fonts, label):
                self.assertIsNotNone(fonts, label)
                self.assertEqual(fonts.get(qn("w:ascii")), "Times New Roman", label)
                self.assertEqual(fonts.get(qn("w:hAnsi")), "Times New Roman", label)
                self.assertEqual(fonts.get(qn("w:eastAsia")), "SimSun", label)
                self.assertIsNone(fonts.get(qn("w:asciiTheme")), label)
                self.assertIsNone(fonts.get(qn("w:hAnsiTheme")), label)
                self.assertIsNone(fonts.get(qn("w:eastAsiaTheme")), label)
                self.assertIsNone(fonts.get(qn("w:cstheme")), label)
                for attr in (
                    qn("w:ascii"), qn("w:hAnsi"), qn("w:cs"), qn("w:eastAsia"),
                ):
                    value = fonts.get(attr)
                    if value is None:
                        continue
                    self.assertIn(
                        value,
                        {"Times New Roman", "SimSun"},
                        f"{label} 出现未允许字体: {value}",
                    )

            for style_name in (
                "Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3",
            ):
                fonts = document.styles[style_name]._element.rPr.rFonts
                assert_bilingual_fonts(fonts, style_name)

            heading_names = {"Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"}
            heading_found = set()
            for paragraph in document.paragraphs:
                style_name = paragraph.style.name if paragraph.style else ""
                if paragraph.text == REPORT_CONFIG["project_name"]:
                    self.assertTrue(paragraph.runs)
                    for run in paragraph.runs:
                        fonts = run._element.rPr.rFonts
                        self.assertEqual(fonts.get(qn("w:ascii")), "SimSun")
                        self.assertEqual(fonts.get(qn("w:hAnsi")), "SimSun")
                        self.assertEqual(fonts.get(qn("w:eastAsia")), "SimSun")
                        self.assertFalse(run.font.italic)
                    heading_found.add(style_name)
                    continue
                for run in paragraph.runs:
                    rpr = run._element.rPr
                    if rpr is None or rpr.rFonts is None:
                        continue
                    assert_bilingual_fonts(
                        rpr.rFonts,
                        f"{style_name}:{run.text[:20]}",
                    )
                if style_name in heading_names and paragraph.text.strip():
                    heading_found.add(style_name)
                    self.assertTrue(paragraph.runs, style_name)
                    for run in paragraph.runs:
                        assert_bilingual_fonts(
                            run._element.rPr.rFonts,
                            f"标题 {style_name}:{paragraph.text}",
                        )
            self.assertGreaterEqual(heading_found, {"Title", "Heading 1", "Heading 2"})

            table_run = next(
                run
                for run in document.tables[0].cell(0, 0).paragraphs[0].runs
                if run.text
            )
            assert_bilingual_fonts(table_run._element.rPr.rFonts, "cover-table")

    def test_custom_project_name_reaches_html_and_word_cover(self):
        project_name = "某集团数据库专项巡检"
        with tempfile.TemporaryDirectory() as td:
            raw = self._make_collection(td, "raw", "db01", sid="orcl", kind="db")
            output = Path(td, "custom.html")
            result = build_report(
                [raw], str(output), project_name=project_name,
            )
            self.assertEqual(result.env_info["project_name"], project_name)
            self.assertIn(project_name, Path(result.html).read_text(encoding="utf-8"))

            document = Document(result.docx)
            paragraph = next(p for p in document.paragraphs if p.text == project_name)
            self.assertTrue(paragraph.runs)
            for run in paragraph.runs:
                fonts = run._element.rPr.rFonts
                self.assertEqual(fonts.get(qn("w:ascii")), "SimSun")
                self.assertEqual(fonts.get(qn("w:hAnsi")), "SimSun")
                self.assertEqual(fonts.get(qn("w:eastAsia")), "SimSun")
                self.assertFalse(run.font.italic)
            footer_text = "\n".join(
                paragraph.text
                for section in document.sections
                for paragraph in section.footer.paragraphs
            )
            self.assertIn(project_name, footer_text)
            self.assertNotIn("Dotease Limited", footer_text)

    def _sample_systems(self):
        return [
            SystemSummary(
                hostname="app01",
                server_ip="10.0.0.1",
                oracle_sid="orcl",
                timestamp="20260821_101500",
                formatted_time="2026-08-21 10:15:00",
                health_score=70,
                overall_status="存在严重问题",
                counts={"OK": 8, "WARN": 1, "CRIT": 1, "TOTAL": 10},
                categories=["host", "db"],
                issues=[
                    IssueRef(
                        "db", "数据库巡检", "表空间使用率", "CRIT", "96%",
                        "扩容表空间", "app01 / orcl", "app01", "orcl",
                    ),
                    IssueRef(
                        "host", "主机巡检", "磁盘使用率-/", "WARN", "85%",
                        "", "app01 / orcl", "app01", "orcl",
                    ),
                ],
                html_name="Oracle巡检报告_orcl_20260821.html",
                label="app01 / orcl",
            ),
            SystemSummary(
                hostname="app02",
                server_ip="10.0.0.2",
                oracle_sid="prod",
                timestamp="20260821_111500",
                formatted_time="2026-08-21 11:15:00",
                health_score=50,
                overall_status="存在严重问题",
                counts={"OK": 7, "WARN": 0, "CRIT": 3, "TOTAL": 10},
                categories=["host", "db"],
                issues=[
                    IssueRef(
                        "db", "数据库巡检", "表空间使用率", "CRIT", "97%",
                        "扩容表空间", "app02 / prod", "app02", "prod",
                    ),
                    IssueRef(
                        "db", "数据库巡检", "采集完整性", "CRIT", "2项失败",
                        "修复采集", "app02 / prod", "app02", "prod",
                    ),
                ],
                html_name="Oracle巡检报告_prod_20260821.html",
                label="app02 / prod",
            ),
        ]

    def test_summary_model_clusters_risk_and_ranking(self):
        model = build_summary_model(self._sample_systems())
        self.assertEqual(model.risk_level, "高")
        self.assertEqual(model.overall_status, "存在严重问题")
        self.assertEqual(model.fleet_score, 60)
        self.assertEqual(model.avg_score, 60)
        self.assertEqual(model.crit_systems, 2)
        self.assertEqual(len(model.clusters), 1)
        self.assertEqual(model.clusters[0].name, "表空间使用率")
        self.assertEqual(len(model.clusters[0].systems), 2)
        self.assertEqual(model.ranked[0].label, "app02 / prod")
        self.assertIn("优先关注", "".join(model.narrative))
        self.assertTrue(any("共性风险" in item for item in model.narrative))

    def test_summary_fleet_score_uses_single_report_weighted_components(self):
        systems = self._sample_systems()
        systems[0].score_earned = 240
        systems[0].score_possible = 300
        systems[1].score_earned = 60
        systems[1].score_possible = 100
        model = build_summary_model(systems)
        self.assertEqual(model.fleet_score, 75)

    def test_summary_high_score_color_is_independent_from_critical_risk(self):
        system = SystemSummary(
            hostname="db01",
            server_ip="10.0.0.1",
            oracle_sid="orcl",
            timestamp="20260827_120000",
            formatted_time="2026-08-27 12:00:00",
            health_score=90,
            overall_status="存在严重问题",
            confidence="完整",
            score_earned=27,
            score_possible=30,
            counts={
                "OK": 9, "WARN": 0, "CRIT": 1,
                "INFO": 0, "UNKNOWN": 0, "TOTAL": 10,
            },
            categories=["db"],
            label="db01 / orcl",
        )
        model = build_summary_model([system])
        self.assertEqual(model.fleet_score, 90)
        self.assertEqual(model.risk_level, "高")

        with tempfile.TemporaryDirectory() as td:
            html_path = Path(td, "summary.html")
            docx_path = Path(td, "summary.docx")
            generate_summary_html(model, str(html_path))
            generate_summary_docx(model, str(docx_path))

            report = html_path.read_text(encoding="utf-8")
            self.assertIn('<div class="score-ring good">90</div>', report)
            self.assertIn('class="risk-pill risk-high"', report)

            document = Document(docx_path)
            score_paragraph = next(
                paragraph for paragraph in document.paragraphs
                if "整体健康评分 / 100" in paragraph.text
            )
            score_run = next(run for run in score_paragraph.runs if run.text == "90")
            risk_run = next(run for run in score_paragraph.runs if "存在严重问题" in run.text)
            self.assertEqual(str(score_run.font.color.rgb), DOCX_COLORS["OK"])
            self.assertEqual(str(risk_run.font.color.rgb), DOCX_COLORS["CRIT"])

    def test_summary_html_and_docx_are_professional_reports(self):
        model = build_summary_model(self._sample_systems())
        project_name = "某集团数据库巡检项目"
        with tempfile.TemporaryDirectory() as td:
            html_path = Path(td, "Oracle巡检汇总摘要_20260821.html")
            docx_path = Path(td, "Oracle巡检汇总摘要_20260821.docx")
            generate_summary_html(model, str(html_path))
            generate_summary_docx(model, str(docx_path), project_name=project_name)
            report = html_path.read_text(encoding="utf-8")
            self.assertIn("Oracle 巡检汇总摘要报告", report)
            self.assertIn("INS-SUM-20260821", report)
            self.assertIn("app01 / orcl", report)
            self.assertIn("app02 / prod", report)
            self.assertIn("表空间使用率", report)
            self.assertIn("出现于 2 套系统", report)
            self.assertIn("签批栏", report)
            self.assertIn("编制", report)
            self.assertNotIn("${OVERVIEW_NAV}", report)
            self.assertNotIn("${REPORT_BODY}", report)

            document = Document(str(docx_path))
            self.assertNotIn("w:updateFields", document.settings._element.xml)
            self.assertNotIn("TOC", document.element.body.xml)
            first_title_index = next(
                index for index, paragraph in enumerate(document.paragraphs)
                if paragraph.style.name == "Title"
            )
            self.assertNotIn(
                'w:type="page"',
                "".join(
                    paragraph._p.xml
                    for paragraph in document.paragraphs[:first_title_index]
                ),
            )
            heads = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style.name in {"Title", "Heading 1", "Heading 2"}
            ]
            self.assertIn("Oracle 巡检汇总摘要报告", heads)
            self.assertIn("1. 执行摘要", heads)
            self.assertIn("2. 系统健康一览", heads)
            self.assertIn("4. 跨系统共性风险", heads)
            self.assertIn("5. 汇总待处理事项", heads)
            self.assertIn("7. 结论与建议", heads)
            self.assertIn("签批栏", heads)
            all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("表空间使用率", all_text)
            self.assertIn("app02 / prod", all_text)
            footer_text = "\n".join(
                paragraph.text
                for section in document.sections
                for paragraph in section.footer.paragraphs
            )
            self.assertIn(project_name, footer_text)

    def test_build_reports_can_skip_summary(self):
        with tempfile.TemporaryDirectory() as td:
            first = self._make_collection(td, "raw_a", "app01", kind="host")
            second = self._make_collection(td, "raw_b", "app02", kind="host")
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            batch = build_reports(
                [first, second], str(out_dir), write_docx=False, write_summary=False
            )
            self.assertIsNone(batch.summary_html)
            self.assertIsNone(batch.summary_docx)
            html_files = sorted(path.name for path in out_dir.glob("*.html"))
            self.assertEqual(
                html_files,
                [
                    "主机巡检报告_app01_20260821.html",
                    "主机巡检报告_app02_20260821.html",
                ],
            )

    def test_single_report_does_not_write_summary(self):
        with tempfile.TemporaryDirectory() as td:
            first = self._make_collection(td, "raw_a", "app01", kind="host")
            out_dir = Path(td, "reports")
            out_dir.mkdir()
            batch = build_reports([first], str(out_dir), write_docx=False)
            self.assertEqual(len(batch.reports), 1)
            self.assertIsNone(batch.summary_html)
            self.assertEqual(list(out_dir.glob("*汇总摘要*")), [])


if __name__ == "__main__":
    unittest.main()
